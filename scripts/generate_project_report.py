"""Generate the comprehensive project status report as a styled DOCX file."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "reports" / "Iraqi_Legal_Chatbot_Comprehensive_Report.docx"

NAVY = "17324D"
BLUE = "21618C"
TEAL = "148F77"
GOLD = "D4AC0D"
RED = "B03A2E"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "E8F6F3"
LIGHT_GOLD = "FCF3CF"
LIGHT_RED = "FDEDEC"
LIGHT_GRAY = "F4F6F7"
MID_GRAY = "7B8A8B"
WHITE = "FFFFFF"
DARK = "1B2631"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_rtl(paragraph, enabled: bool = True) -> None:
    properties = paragraph._p.get_or_add_pPr()
    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        properties.append(bidi)
    bidi.set(qn("w:val"), "1" if enabled else "0")


def set_run_font(run, name="Aptos", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MID_GRAY)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and select Update Field to refresh this contents page."
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Title", 30, NAVY),
        ("Subtitle", 14, BLUE),
        ("Heading 1", 19, NAVY),
        ("Heading 2", 14, BLUE),
        ("Heading 3", 11.5, TEAL),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(7.0))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    set_cell_shading(left, NAVY)
    set_cell_shading(right, NAVY)
    left.text = "IRAQI LEGAL RAG CHATBOT"
    right.text = "PROJECT REPORT"
    for cell in (left, right):
        set_cell_margins(cell, 70, 100, 70, 100)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=8.5, bold=True, color=WHITE)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.add_run("Confidential project status • ")
    add_page_number(paragraph)


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(45)

    label_table = document.add_table(rows=1, cols=1)
    label_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = label_table.cell(0, 0)
    set_cell_shading(cell, TEAL)
    set_cell_margins(cell, 75, 130, 75, 130)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run("COMPREHENSIVE TECHNICAL REPORT")
    set_run_font(run, size=10, bold=True, color=WHITE)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(22)
    title.paragraph_format.space_after = Pt(8)
    title.add_run("Iraqi Legal\nRAG Chatbot")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Implementation history, current PostgreSQL architecture, "
        "retired OCR/document pipelines, gaps, risks, and next steps"
    )

    document.add_paragraph("")
    architecture = document.add_table(rows=1, cols=4)
    architecture.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ("PostgreSQL", "Ollama", "ChromaDB", "API & UI")
    colors = (NAVY, BLUE, TEAL, GOLD)
    for cell, label, color in zip(architecture.rows[0].cells, labels, colors):
        set_cell_shading(cell, color)
        set_cell_margins(cell, 170, 60, 170, 60)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        set_run_font(run, size=10, bold=True, color=WHITE)

    document.add_paragraph("")
    meta = document.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    values = (
        ("Report date", date.today().isoformat()),
        ("Repository", r"C:\Users\lenovo\chatbot"),
        ("Current source of truth", "PostgreSQL public.iraqi_laws"),
        ("Current status", "Operational index built from 54 database records"),
    )
    for row, (key, value) in zip(meta.rows, values):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_shading(row.cells[1], LIGHT_GRAY)
        row.cells[0].text = key
        row.cells[1].text = value
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, bold=True, color=NAVY)
        for cell in row.cells:
            set_cell_margins(cell)

    document.add_paragraph("")
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    run = note.add_run(
        "Security note: credentials are intentionally excluded. "
        "The previously exposed database password should be rotated."
    )
    set_run_font(run, size=9.5, bold=True, color=RED)
    document.add_page_break()


def add_heading(document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_body(document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True, color=NAVY)
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def add_bullets(document, items, level=0) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
        paragraph.add_run(item)


def add_numbered(document, items) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_status_table(document, rows) -> None:
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    headers = ("Area", "Status", "Evidence / result")
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        cell.text = label
        for run in cell.paragraphs[0].runs:
            set_run_font(run, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    status_colors = {
        "Completed": LIGHT_TEAL,
        "Removed": LIGHT_RED,
        "Not tried": LIGHT_GOLD,
        "Partial": LIGHT_BLUE,
        "Current": LIGHT_TEAL,
    }
    for area, status, evidence in rows:
        cells = table.add_row().cells
        cells[0].text = area
        cells[1].text = status
        cells[2].text = evidence
        set_cell_shading(cells[1], status_colors.get(status, LIGHT_GRAY))
        for run in cells[1].paragraphs[0].runs:
            set_run_font(run, bold=True, color=DARK)
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_code_block(document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "202B33")
    set_cell_margins(cell, 130, 160, 130, 160)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(text.splitlines()):
        if index:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        set_run_font(run, name="Cascadia Mono", size=8.5, color="E8F0F2")


def add_callout(document, title: str, text: str, color=BLUE, fill=LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=2)
    table.columns[0].width = Inches(0.08)
    table.columns[1].width = Inches(6.7)
    set_cell_shading(table.cell(0, 0), color)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 1), 120, 150, 120, 150)
    paragraph = table.cell(0, 1).paragraphs[0]
    run = paragraph.add_run(f"{title}\n")
    set_run_font(run, bold=True, color=color)
    paragraph.add_run(text)


def build_report() -> Document:
    document = Document()
    configure_document(document)
    add_cover(document)

    add_heading(document, "Contents")
    add_toc(document.add_paragraph())
    document.add_page_break()

    add_heading(document, "1. Executive Summary")
    add_body(
        document,
        "The project is an offline Arabic legal retrieval-augmented generation "
        "system. It evolved through PDF extraction, scanned-PDF OCR, DOCX/TXT "
        "batch ingestion, filesystem CRUD, and finally a PostgreSQL-only "
        "knowledge architecture. The current index was successfully rebuilt "
        "from 54 rows in public.iraqi_laws."
    )
    add_callout(
        document,
        "Current decision",
        "PostgreSQL is the sole durable legal source. ChromaDB is a generated "
        "retrieval index; qwen2.5:7b is not fine-tuned and answers using RAG.",
        TEAL,
        LIGHT_TEAL,
    )
    add_status_table(
        document,
        (
            ("PostgreSQL connection", "Completed", "Connected as chatbot_app"),
            ("PostgreSQL ingestion", "Completed", "54 rows → 54 chunks"),
            ("Embedding/index build", "Completed", "bge-m3 → ChromaDB"),
            ("PDF/DOCX/TXT ingestion", "Removed", "No longer a supported path"),
            ("Scanned-PDF OCR", "Removed", "Historically built and tested"),
            ("Model fine-tuning", "Not tried", "Blocked by data and 2 GB VRAM"),
            ("Answer-quality benchmark", "Not tried", "No reviewed test set"),
        ),
    )

    add_heading(document, "2. Current Architecture")
    add_code_block(
        document,
        "PostgreSQL public.iraqi_laws\n"
        "    ↓ legal metadata mapping\n"
        "54 normalized legal records\n"
        "    ↓ Iraqi legal chunking\n"
        "54 inspectable chunks\n"
        "    ↓ Ollama bge-m3 embeddings\n"
        "ChromaDB iraqi_legal_documents\n"
        "    ↓ semantic + BM25 reranking\n"
        "qwen2.5:7b grounded generation\n"
        "    ↓ citation validation\n"
        "Browser / API / Streamlit / terminal"
    )
    add_heading(document, "2.1 Current data mapping", level=2)
    add_bullets(
        document,
        (
            "id → stable PostgreSQL source identity",
            "classification → document type",
            "law_number and law_year → full legal reference",
            "article_number → article locator",
            "law_name → document title",
            "summary → retrievable legal content",
        ),
    )
    add_heading(document, "2.2 Confirmed runtime evidence", level=2)
    add_bullets(
        document,
        (
            "Database laws_iraqi connected successfully as chatbot_app.",
            "All 54 rows had non-null summaries.",
            "Four embedding batches completed in approximately 22 seconds.",
            "One Chroma batch stored all 54 chunks.",
            "The iraqi_legal_documents collection reported 54 indexed chunks.",
            "The PostgreSQL-only unit suite passed 21 tests.",
        ),
    )

    add_heading(document, "3. Project Evolution Timeline")
    timeline = (
        ("Phase 1", "PDF corpus and Arabic-aware retrieval", "Completed historically"),
        ("Phase 2", "Postman-compatible API and file CRUD", "Completed historically"),
        ("Phase 3", "DOCX/TXT batch ingestion and metadata", "Completed historically"),
        ("Phase 4", "Scanned-PDF OCR and Arabic optimization", "Completed historically"),
        ("Phase 5", "PostgreSQL connection and row mapping", "Current"),
        ("Phase 6", "PostgreSQL-only simplification", "Current"),
        ("Phase 7", "Model fine-tuning", "Not performed"),
    )
    add_status_table(document, timeline)

    add_heading(document, "4. PostgreSQL Integration")
    add_body(
        document,
        "The existing non-empty laws_iraqi database was inspected through a "
        "schema-only export. It contains public.iraqi_laws with an integer "
        "primary key and structured legal columns."
    )
    add_heading(document, "4.1 Completed", level=2)
    add_bullets(
        document,
        (
            "Created and validated the chatbot_app application role.",
            "Granted database, schema, table, and sequence access.",
            "Added SQLAlchemy, Psycopg, and dotenv configuration.",
            "Added pooled connections with pre-ping health checks.",
            "Mapped each row into citation-preserving retrieval metadata.",
            "Rebuilt chunks, citation registry, embeddings, and Chroma index.",
            "Removed supported filesystem ingestion and document CRUD routes.",
        ),
    )
    add_heading(document, "4.2 Not yet implemented", level=2)
    add_bullets(
        document,
        (
            "Incremental indexing after INSERT, UPDATE, or DELETE.",
            "Database triggers, notifications, or a durable indexing job queue.",
            "PostgreSQL CRUD endpoints.",
            "Alembic migration baseline for the existing schema.",
            "PostgreSQL status in the /health response.",
            "Automated backup and restore procedures.",
            "pgvector comparison or migration.",
        ),
    )

    add_heading(document, "5. Historical PDF Pipeline")
    add_body(
        document,
        "The earliest corpus used text-layer PDFs. pdfplumber extracted page "
        "text, which was cleaned, chunked, embedded, and stored in ChromaDB."
    )
    add_heading(document, "5.1 What worked", level=2)
    add_bullets(
        document,
        (
            "Multi-document indexing into ChromaDB.",
            "Inspectable chunks and source/page metadata.",
            "Arabic-aware semantic retrieval.",
            "Source provenance in answers and snippets.",
        ),
    )
    add_heading(document, "5.2 Problems found", level=2)
    add_bullets(
        document,
        (
            "RTL extraction could reverse multi-digit law numbers and years.",
            "The corpus contained roughly 201 chunks across three PDFs.",
            "Several sources resembled explanatory summaries, not primary statutes.",
            "Weak or corrupt source text placed a hard ceiling on answer accuracy.",
        ),
    )
    add_callout(
        document,
        "Current status",
        "PDF loading, dependencies, tests, and operator documentation were "
        "removed after PostgreSQL became the only supported source.",
        RED,
        LIGHT_RED,
    )

    add_heading(document, "6. Historical Scanned-PDF OCR")
    add_body(
        document,
        "OCR was genuinely implemented and tested; it was not merely proposed. "
        "The loader attempted embedded text first and used Tesseract when a "
        "page contained no usable text."
    )
    add_heading(document, "6.1 Implemented OCR pipeline", level=2)
    add_numbered(
        document,
        (
            "Render scanned pages with pypdfium2.",
            "Use 400 DPI to preserve Arabic character detail.",
            "Apply grayscale, autocontrast, sharpening, and thresholding.",
            "Run Tesseract with page segmentation modes 6, 4, and 11.",
            "Score candidates for Arabic characters and legal terminology.",
            "Penalize Latin/noise-heavy output.",
            "Keep the strongest candidate and record OCR provenance.",
        ),
    )
    add_heading(document, "6.2 Dependencies and environment", level=2)
    add_bullets(
        document,
        (
            "pytesseract was the Python bridge.",
            "The Windows Tesseract executable was separately required.",
            "Arabic ara.traineddata had to appear in tesseract --list-langs.",
            "pdfplumber handled text-layer pages.",
            "pypdfium2 rendered scanned pages.",
            "Pillow provided preprocessing operations.",
        ),
    )
    add_heading(document, "6.3 Confirmed OCR failure mode", level=2)
    add_body(
        document,
        "One scanned law produced a corrupted number/year containing Latin "
        "noise. Inspection proved that the corruption was already present in "
        "data/chunks.jsonl. Retrieval and generation could not reconstruct a "
        "legal reference that OCR had stored incorrectly."
    )
    add_callout(
        document,
        "Key lesson",
        "When a legal answer contains the wrong law number or year, inspect the "
        "indexed chunk first. If the chunk is wrong, improve the scan/OCR or "
        "replace the source; prompt tuning cannot repair missing evidence.",
        GOLD,
        LIGHT_GOLD,
    )
    add_heading(document, "6.4 Current status", level=2)
    add_body(
        document,
        "OCR code, tests, dependencies, and documentation were removed in the "
        "PostgreSQL-only refactor. Existing files were left untouched. The "
        "current installation state of Tesseract was not revalidated."
    )

    add_heading(document, "7. Historical DOCX and TXT Pipeline")
    add_bullets(
        document,
        (
            "Batch loading from data/legal_documents.",
            "DOCX headings, paragraphs, and tables preserved in order.",
            "Word title and core properties propagated as metadata.",
            "TXT and legacy page separators supported.",
            "Article, section, title, source, and document type preserved.",
            "Filesystem watcher supported periodic rebuilding.",
            "Text document CRUD synchronized source files and Chroma.",
        ),
    )
    add_body(
        document,
        "The loader, watcher, CRUD module, tests, python-docx dependency, and "
        "DOCX migration guide were later removed. The project no longer reads "
        "PDF, DOCX, or TXT knowledge files."
    )

    add_heading(document, "8. Retrieval and Answer Quality Work")
    add_heading(document, "8.1 Completed improvements", level=2)
    add_bullets(
        document,
        (
            "Arabic letter normalization.",
            "Arabic and Western digit normalization.",
            "Alternate matching for reversed RTL digit sequences.",
            "Semantic candidate retrieval through bge-m3 and Chroma.",
            "BM25-style lexical relevance scoring.",
            "Higher weighting for exact law numbers and years.",
            "Candidate expansion before final selection.",
            "Near-duplicate suppression.",
            "Shared citation validation across terminal and Streamlit.",
            "Non-blocking warnings when citation verification remains weak.",
        ),
    )
    add_heading(document, "8.2 Not yet measured", level=2)
    add_bullets(
        document,
        (
            "Recall@k on a reviewed legal question set.",
            "Mean reciprocal rank.",
            "Citation precision and completeness.",
            "Answer correctness scored by a legal reviewer.",
            "Regression testing after database changes.",
        ),
    )

    add_heading(document, "9. API and User Interfaces")
    add_status_table(
        document,
        (
            ("GET /health", "Current", "Lightweight process health"),
            ("POST /ask", "Current", "Grounded answer API"),
            ("Browser frontend", "Current", "Served by Starlette"),
            ("Streamlit", "Current", "Shared answer logic"),
            ("Terminal", "Current", "Shared answer logic"),
            ("Filesystem /documents CRUD", "Removed", "No longer routed"),
            ("Frontend upload manager", "Partial", "UI assets still remain"),
        ),
    )
    add_callout(
        document,
        "Known inconsistency",
        "The browser still contains historical upload-manager components even "
        "though the backend document endpoints were removed. This UI should be "
        "removed or redesigned as PostgreSQL management.",
        RED,
        LIGHT_RED,
    )

    add_heading(document, "10. Model Training Assessment")
    add_body(
        document,
        "The qwen2.5:7b model has not been fine-tuned. The current system uses "
        "retrieval-augmented generation, which supplies relevant PostgreSQL "
        "records to the model at answer time."
    )
    add_status_table(
        document,
        (
            ("GPU", "Partial", "NVIDIA MX550 with 2 GB VRAM"),
            ("Training corpus", "Partial", "54 summaries only"),
            ("PyTorch/Transformers", "Not tried", "Not installed"),
            ("PEFT/TRL/Datasets", "Not tried", "Not installed"),
            ("LoRA/QLoRA run", "Not tried", "Hardware insufficient for 7B"),
            ("Reviewed instruction data", "Not tried", "No training set"),
            ("GGUF export/Ollama import", "Not tried", "No trained adapter"),
        ),
    )
    add_body(
        document,
        "Fine-tuning should teach answer style, refusal behavior, and citation "
        "format—not serve as the legal database. Current legal facts should "
        "continue to come from PostgreSQL retrieval."
    )

    add_heading(document, "11. Validation Performed")
    add_bullets(
        document,
        (
            "Historical file-ingestion suite: 36 tests passed before removal.",
            "Current PostgreSQL-only suite: 21 tests passed.",
            "Python compile validation passed.",
            "Git whitespace validation passed.",
            "PostgreSQL connection was tested live.",
            "Full PostgreSQL → embedding → Chroma rebuild succeeded.",
            "A live reviewed set of /ask answers has not yet been completed.",
        ),
    )

    add_heading(document, "12. Tried vs. Not Tried")
    add_status_table(
        document,
        (
            ("PDF text extraction", "Completed", "Later removed"),
            ("Scanned-PDF OCR", "Completed", "Later removed"),
            ("Arabic OCR tuning", "Completed", "400 DPI and multi-pass"),
            ("DOCX/TXT ingestion", "Completed", "Later removed"),
            ("Filesystem CRUD", "Completed", "Later removed"),
            ("PostgreSQL ingestion", "Current", "54 rows indexed"),
            ("Arabic retrieval reranking", "Current", "Tests passing"),
            ("Structured citations", "Current", "Registry-backed"),
            ("Incremental DB synchronization", "Not tried", "Manual rebuild"),
            ("PostgreSQL CRUD API", "Not tried", "No routes"),
            ("pgvector", "Not tried", "Chroma retained"),
            ("Fine-tuning", "Not tried", "Hardware/data blocker"),
            ("Legal expert evaluation", "Not tried", "No benchmark"),
            ("Production deployment", "Not tried", "Local only"),
            ("API authentication", "Not tried", "No security layer"),
        ),
    )

    add_heading(document, "13. Current Risks")
    risks = (
        ("Critical", "Exposed database password", "Rotate the chatbot_app password."),
        ("High", "Summaries may not be authoritative law text", "Store full official articles."),
        ("High", "Chroma becomes stale after DB changes", "Add queued incremental indexing."),
        ("High", "No reviewed accuracy benchmark", "Build and score a legal test set."),
        ("Medium", "Upload UI conflicts with PostgreSQL-only backend", "Remove or redesign it."),
        ("Medium", "No amendment/repeal modeling", "Add status and effective dates."),
        ("Medium", "Uncommitted worktree", "Review and commit the migration cleanly."),
        ("Low", "laws_schema.sql remains untracked", "Archive or ignore intentionally."),
    )
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("Severity", "Risk", "Action")):
        set_cell_shading(cell, NAVY)
        cell.text = text
        for run in cell.paragraphs[0].runs:
            set_run_font(run, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    fills = {"Critical": LIGHT_RED, "High": LIGHT_GOLD, "Medium": LIGHT_BLUE, "Low": LIGHT_GRAY}
    for severity, risk, action in risks:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = severity, risk, action
        set_cell_shading(cells[0], fills[severity])
        for run in cells[0].paragraphs[0].runs:
            set_run_font(run, bold=True)
        for cell in cells:
            set_cell_margins(cell)

    add_heading(document, "14. Recommended Roadmap")
    add_numbered(
        document,
        (
            "Rotate the exposed PostgreSQL password and update .env.",
            "Remove or redesign the historical browser upload interface.",
            "Create 20–50 reviewed legal questions with expected sources.",
            "Measure retrieval recall and citation accuracy before prompt changes.",
            "Add automatic PostgreSQL-to-Chroma synchronization.",
            "Store full authoritative article text instead of summaries alone.",
            "Model amendments, repeals, effective dates, and law relationships.",
            "Add PostgreSQL CRUD endpoints only if remote management is required.",
            "Add API authentication, audit logs, backups, and recovery testing.",
            "Consider fine-tuning only after obtaining reviewed training examples "
            "and suitable external GPU hardware.",
        ),
    )

    add_heading(document, "15. Operator Runbook")
    add_heading(document, "15.1 Validate the database", level=2)
    add_code_block(document, r"python -m app.database")
    add_heading(document, "15.2 Rebuild the retrieval index", level=2)
    add_code_block(document, r"python ingest_legal_documents.py")
    add_heading(document, "15.3 Start the browser/API", level=2)
    add_code_block(
        document,
        r"python -m uvicorn app.api:app --host 127.0.0.1 --port 8000",
    )
    add_heading(document, "15.4 Open the application", level=2)
    add_code_block(document, "http://127.0.0.1:8000/")

    add_heading(document, "Appendix A — Current Key Files")
    add_status_table(
        document,
        (
            ("app/database.py", "Current", "Connection and SQL query boundary"),
            ("app/postgres_laws.py", "Current", "Database row mapping"),
            ("app/legal_document.py", "Current", "Structured legal data model"),
            ("app/build_index.py", "Current", "Embedding and Chroma rebuild"),
            ("app/search_index.py", "Current", "Hybrid retrieval/reranking"),
            ("app/rag_answer.py", "Current", "Grounded generation and vetting"),
            ("app/api.py", "Current", "Browser, /health, and /ask"),
            ("data/chunks.jsonl", "Current", "Inspectable generated chunks"),
            ("data/citation_registry.json", "Current", "Generated citations"),
            ("laws_schema.sql", "Partial", "Untracked schema snapshot"),
        ),
    )

    add_heading(document, "Appendix B — Removed Components")
    add_bullets(
        document,
        (
            "app/document_loaders.py",
            "app/document_classifier.py",
            "app/document_store.py",
            "watch_legal_documents.py",
            "docs/docx-ingestion-migration.md",
            "PDF/DOCX/TXT/OCR/classification/document-store tests",
            "pdfplumber, pypdfium2, pytesseract, Pillow, and python-docx runtime dependencies",
        ),
    )

    add_heading(document, "Appendix C — Report Scope")
    add_body(
        document,
        "This report combines current repository inspection, successful runtime "
        "outputs from the PostgreSQL migration, and historical implementation "
        "records for ingestion, OCR, retrieval, API, and citation work. Items "
        "marked not tried were not represented as completed."
    )
    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_report()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
