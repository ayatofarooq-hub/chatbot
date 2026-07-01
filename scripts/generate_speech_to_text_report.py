"""Generate the Arabic Word report for the Jalssa speech-to-text feature."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "تقرير-خاصية-تحويل-الكلام-إلى-نص-منصة-جلسة.docx"
LOGO = ROOT / "frontend" / "images" / "logo-icon.png"

GREEN = "145A38"
GREEN_DARK = "0D472D"
GREEN_SOFT = "E8F5ED"
GOLD = "C8A008"
INK = "171D19"
MUTED = "637069"
BORDER = "E2E8E4"
WHITE = "FFFFFF"


def set_cell_shading(cell, color):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size="6"):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")
    return paragraph


def style_run(run, size=11, bold=False, color=INK, font="Arial"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:cs"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return run


def add_text(document, text, size=11, bold=False, color=INK, before=0, after=7):
    paragraph = rtl(document.add_paragraph())
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    style_run(paragraph.add_run(text), size=size, bold=bold, color=color)
    return paragraph


def add_heading(document, number, title):
    paragraph = rtl(document.add_paragraph())
    paragraph.paragraph_format.space_before = Pt(15)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    style_run(paragraph.add_run(f"{number}  {title}"), size=17, bold=True, color=GREEN)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), GOLD)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)
    return paragraph


def add_bullet(document, text, color=INK):
    paragraph = rtl(document.add_paragraph())
    paragraph.paragraph_format.left_indent = Cm(0.3)
    paragraph.paragraph_format.right_indent = Cm(0.3)
    paragraph.paragraph_format.space_after = Pt(4)
    style_run(paragraph.add_run("◆  "), size=8, bold=True, color=GOLD)
    style_run(paragraph.add_run(text), size=11, color=color)
    return paragraph


def add_info_box(document, title, body, color=GREEN_SOFT):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.2)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, color)
    set_cell_border(cell, GREEN, "8")
    p = rtl(cell.paragraphs[0])
    p.paragraph_format.space_after = Pt(4)
    style_run(p.add_run(title), size=12, bold=True, color=GREEN)
    p2 = rtl(cell.add_paragraph())
    p2.paragraph_format.space_after = Pt(3)
    p2.paragraph_format.line_spacing = 1.2
    style_run(p2.add_run(body), size=10.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_matrix(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, GREEN)
        set_cell_border(cell, WHITE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Cm(widths[index])
        p = rtl(cell.paragraphs[0])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(header), size=10, bold=True, color=WHITE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            if widths:
                cell.width = Cm(widths[index])
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else "F7FAF8")
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = rtl(cell.paragraphs[0])
            p.paragraph_format.space_after = Pt(2)
            style_run(p.add_run(str(value)), size=9.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_flow(document):
    steps = [
        ("1", "الميكروفون", "التقاط السؤال"),
        ("2", "Web Audio PCM", "إنشاء WAV مباشر"),
        ("3", "/transcribe", "التحقق والاستقبال"),
        ("4", "Whisper", "تحويل عربي محلي"),
        ("5", "حقل السؤال", "المراجعة ثم الإرسال"),
    ]
    table = document.add_table(rows=1, cols=len(steps))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (number, title, note) in enumerate(reversed(steps)):
        cell = table.cell(0, index)
        cell.width = Cm(3.2)
        set_cell_shading(cell, GREEN_SOFT if index % 2 == 0 else "F8FBF9")
        set_cell_border(cell, GREEN, "8")
        p = rtl(cell.paragraphs[0])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(number), size=17, bold=True, color=GOLD)
        p2 = rtl(cell.add_paragraph())
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p2.add_run(title), size=10, bold=True, color=GREEN_DARK)
        p3 = rtl(cell.add_paragraph())
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p3.add_run(note), size=8, color=MUTED)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")

    header = section.header
    header.is_linked_to_previous = False
    hp = rtl(header.paragraphs[0])
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(hp.add_run("منصة جلسة  |  التقرير التقني"), size=9, bold=True, color=GREEN)

    footer = section.footer
    fp = rtl(footer.paragraphs[0])
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(fp.add_run("خاصية تحويل الكلام إلى نص  •  وثيقة داخلية"), size=8, color=MUTED)
    style_run(fp.add_run("    |    "), size=8, color=GOLD)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    fp._p.append(field)


def add_cover(document):
    for _ in range(2):
        document.add_paragraph()
    if LOGO.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Cm(2.8))

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    style_run(p.add_run("منصة جلسة"), size=28, bold=True, color=GREEN_DARK)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(15)
    style_run(p.add_run("تقرير خاصية تحويل الكلام إلى نص"), size=25, bold=True, color=GREEN)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    style_run(p.add_run("Speech-to-Text"), size=16, bold=True, color=GOLD)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    style_run(
        p.add_run("التصميم، المعمارية، تنفيذ الواجهة الأمامية والخلفية، الأمان والاختبارات"),
        size=12,
        color=MUTED,
    )

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, GREEN)
    set_cell_border(cell, GOLD, "14")
    p = rtl(cell.paragraphs[0])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run("تنفيذ محلي • دعم اللغة العربية • خصوصية التسجيلات"), size=12, bold=True, color=WHITE)

    for _ in range(3):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run("الإصدار 1.1  |  يوليو 2026"), size=10, color=MUTED)
    document.add_page_break()


def build_report():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)

    add_heading(document, "01", "الملخص التنفيذي")
    add_text(
        document,
        "تم تطوير خاصية تحويل الكلام إلى نص داخل منصة جلسة لتمكين المستخدم من إملاء السؤال القانوني باللغة العربية، "
        "ثم مراجعة النص الناتج قبل إرساله إلى المساعد القانوني. تعمل عملية التحويل محلياً بواسطة faster-whisper، "
        "ولا تعتمد على خدمات سحابية بعد تنزيل النموذج لأول مرة.",
    )
    add_info_box(
        document,
        "النتيجة الأساسية",
        "مسار مستقل وآمن: تسجيل من المتصفح ← رفع مؤقت ← تحويل عربي محلي ← إعادة النص إلى حقل السؤال. "
        "لا يُرسل السؤال تلقائياً ولا يتم الاحتفاظ بالتسجيل الصوتي.",
    )
    add_matrix(
        document,
        ["المؤشر", "القيمة الحالية"],
        [
            ("النموذج", "faster-whisper small"),
            ("لغة التحويل", "العربية ar"),
            ("مكان المعالجة", "خادم منصة جلسة محلياً"),
            ("مدة التسجيل", "من ثانيتين إلى 60 ثانية"),
            ("حجم الملف", "10 MB كحد أقصى"),
            ("الاختبارات", "13 اختباراً ناجحاً"),
        ],
        [8, 8],
    )

    add_heading(document, "02", "تجربة المستخدم")
    add_bullet(document, "يضغط المستخدم على زر الميكروفون لبدء التسجيل.")
    add_bullet(document, "يتحول الزر إلى مربع أحمر مع رسالة تؤكد بدء التسجيل.")
    add_bullet(document, "يضغط المستخدم مرة أخرى لإيقاف التسجيل، أو يتوقف تلقائياً بعد 60 ثانية.")
    add_bullet(document, "تظهر حالة التحويل، ثم يُضاف النص الناتج إلى حقل السؤال.")
    add_bullet(document, "يراجع المستخدم المصطلحات القانونية والأرقام قبل الضغط على إرسال.")
    add_info_box(
        document,
        "قرار تصميمي",
        "لا يتم إرسال السؤال تلقائياً بعد التحويل. المراجعة البشرية مهمة لأن أسماء القوانين وأرقام المواد "
        "قد تتأثر باللهجة، الضوضاء، أو جودة الميكروفون.",
        "FFF8DB",
    )

    add_heading(document, "03", "المعمارية ومسار البيانات")
    add_flow(document)
    add_text(
        document,
        "Speech-to-Text منفصل عن نموذج الإجابة Ollama. Whisper مسؤول فقط عن تحويل الصوت إلى نص، "
        "بينما ينتقل النص بعد موافقة المستخدم إلى مسار /ask، ثم البحث القانوني ونظام RAG وإنتاج الإجابة.",
    )

    add_heading(document, "04", "تنفيذ الواجهة الأمامية")
    add_matrix(
        document,
        ["المكوّن", "المسؤولية", "التقنية"],
        [
            ("زر الصوت", "بدء وإيقاف التسجيل وإظهار الحالة", "HTML + CSS"),
            ("الوصول للميكروفون", "طلب الإذن والتقاط المسار الصوتي", "getUserMedia"),
            ("التسجيل الأساسي", "التقاط PCM وإنشاء WAV داخل المتصفح", "Web Audio API"),
            ("اختيار الجهاز", "عرض أجهزة الإدخال وحفظ deviceId", "enumerateDevices"),
            ("رفع التسجيل", "إرسال multipart/form-data", "Fetch + FormData"),
            ("إدخال النص", "إضافة الناتج إلى مربع السؤال", "DOM events"),
        ],
        [3.2, 7.1, 5.7],
    )
    add_text(document, "آلية الالتقاط النهائية:", bold=True, color=GREEN_DARK)
    for item in (
        "التقاط عينات PCM مباشرة من Web Audio API.",
        "إنشاء ملف WAV أحادي القناة بصيغة PCM 16-bit داخل المتصفح.",
        "تجميع التسجيل في الذاكرة ثم تحرير مسار الميكروفون فور الإيقاف.",
        "استخدام MediaRecorder فقط كمسار احتياطي عند غياب Web Audio.",
    ):
        add_bullet(document, item)
    add_text(document, "صيغة الإرسال الأساسية:", bold=True, color=GREEN_DARK)
    add_bullet(document, "WAV/PCM، مع WebM أو OGG أو MP4 كصيغ احتياطية عند الحاجة.")
    add_text(document, "فحص جودة الإدخال:", bold=True, color=GREEN_DARK)
    add_bullet(document, "رفض التسجيل الأقل من ثانيتين.")
    add_bullet(document, "قياس مدة التسجيل وRMS وPeak على الخادم لتشخيص الملفات الصامتة دون الاحتفاظ بها.")
    add_bullet(document, "إظهار رسائل مختلفة لرفض الإذن، انشغال الجهاز، أو عدم دعم القيود.")

    document.add_page_break()
    add_heading(document, "05", "تنفيذ الواجهة الخلفية")
    add_text(
        document,
        "تمت إضافة POST /transcribe إلى خادم Starlette. يستقبل المسار ملفاً صوتياً في الحقل audio، "
        "ويتحقق من الطلب قبل تمريره إلى محرك التحويل.",
    )
    add_matrix(
        document,
        ["التحقق", "السلوك"],
        [
            ("صيغة الطلب", "multipart/form-data"),
            ("الحقل المطلوب", "audio"),
            ("الأنواع", "MP4، MPEG، OGG، WAV، WebM"),
            ("الملف الفارغ", "رفض 422"),
            ("النوع غير المدعوم", "رفض 415"),
            ("أكبر من 10 MB", "رفض 413"),
            ("خطأ تثبيت المحرك", "إرجاع 503"),
            ("فشل عام", "إرجاع 500"),
        ],
        [6, 10],
    )
    add_text(
        document,
        "يُحفظ الملف باستخدام NamedTemporaryFile، وتُنفذ عملية Whisper داخل run_in_threadpool حتى لا تحجب "
        "event loop. يُحذف الملف في finally سواء نجحت العملية أو فشلت.",
    )
    add_info_box(
        document,
        "عقد الاستجابة",
        'النجاح: {"text": "النص المحول", "language": "ar"} — أما الخطأ فيرجع داخل الحقل detail مع رمز HTTP مناسب.',
    )

    add_heading(document, "06", "محرك Whisper وإدارة النموذج")
    add_text(
        document,
        "تم استخدام faster-whisper 1.2.1، وهو تنفيذ محسن من Whisper يعتمد على CTranslate2 وPyAV. "
        "النموذج الافتراضي small ويعمل على CPU باستخدام int8 لتقليل استهلاك الذاكرة.",
    )
    add_matrix(
        document,
        ["الإعداد", "القيمة", "الغرض"],
        [
            ("language", "ar", "توجيه التحويل إلى العربية"),
            ("vad_filter", "True", "إزالة الصمت في المحاولة الأولى"),
            ("beam_size", "5", "تحسين اختيار النص"),
            ("condition_on_previous_text", "False", "تقليل التكرار والهلوسة"),
            ("المحاولة الاحتياطية", "VAD=False", "معالجة الحالات التي يصنف فيها الكلام كصمت"),
        ],
        [4.2, 3.2, 8.6],
    )
    add_text(document, "متغيرات البيئة المدعومة:", bold=True, color=GREEN_DARK)
    add_bullet(document, "WHISPER_MODEL لتغيير اسم النموذج أو تحديد مسار محلي.")
    add_bullet(document, "WHISPER_DEVICE لاختيار cpu أو cuda.")
    add_bullet(document, "WHISPER_COMPUTE_TYPE لاختيار int8 أو float16 حسب الجهاز.")
    add_text(
        document,
        "يُحمّل النموذج عند أول طلب فقط، ثم يبقى في الذاكرة. يحمي threading.Lock من تحميل نسخ متعددة "
        "عند وصول طلبات متزامنة.",
    )

    add_heading(document, "07", "الخصوصية والأمان")
    add_bullet(document, "المعالجة محلية ولا تُرسل التسجيلات إلى Google أو OpenAI أو مزود سحابي.")
    add_bullet(document, "يُحذف الملف المؤقت فور انتهاء التحويل.")
    add_bullet(document, "لا يُرسل النص تلقائياً إلى نظام الإجابة.")
    add_bullet(document, "يجب استخدام HTTPS في النشر الخارجي لأن المتصفحات تقيد الميكروفون خارج السياق الآمن.")
    add_bullet(document, "المسار الحالي لا يتطلب جلسة مسؤول لأنه لا يصل إلى الوثائق أو الإعدادات.")
    add_info_box(
        document,
        "متطلبات الإنتاج",
        "إضافة مصادقة مناسبة للمستخدمين، rate limiting، حد للتزامن، سجلات تشغيل دون صوت، سياسة خصوصية واضحة، "
        "ومراقبة استهلاك CPU/GPU.",
        "FFF8DB",
    )

    add_heading(document, "08", "الاختبارات والتحقق")
    add_matrix(
        document,
        ["نوع الاختبار", "النتيجة"],
        [
            ("قبول ملف صوت صحيح وإعادة النص", "ناجح"),
            ("رفض النوع غير المدعوم", "ناجح"),
            ("إعادة المحاولة دون VAD", "ناجح"),
            ("رسالة عربية عند غياب الكلام", "ناجح"),
            ("اختبار بملف صوت فعلي عبر /transcribe", "ناجح"),
            ("Python compileall وgit diff --check", "ناجح"),
        ],
        [11, 5],
    )
    add_info_box(
        document,
        "حالة التحقق",
        "تم تنفيذ 13 اختباراً آلياً بنجاح، واختبار WAV فعلي من طرف إلى طرف، "
        "كما تم تأكيد نجاح التسجيل والتحويل من واجهة منصة جلسة على الجهاز المستهدف.",
    )

    add_heading(document, "09", "ملفات التنفيذ")
    add_matrix(
        document,
        ["الملف", "الدور"],
        [
            ("frontend/index.html", "زر الميكروفون وربط عناصر الواجهة"),
            ("frontend/app.js", "اختيار الميكروفون، التقاط PCM، إنشاء WAV، الرفع وإدخال النص"),
            ("frontend/styles.css", "حالات التسجيل والتحويل والحركة البصرية"),
            ("app/api.py", "مسار /transcribe والتحقق وإدارة الملف المؤقت"),
            ("app/speech_to_text.py", "تحميل Whisper وتنفيذ التحويل والمحاولة الاحتياطية"),
            ("requirements.txt", "تثبيت faster-whisper"),
            ("tests/test_api.py", "اختبارات عقد API"),
            ("tests/test_speech_to_text.py", "اختبارات منطق Whisper وVAD"),
        ],
        [7, 9],
    )

    add_heading(document, "10", "التشغيل والصيانة")
    add_text(document, "تشغيل المنصة محلياً:", bold=True, color=GREEN_DARK)
    add_info_box(
        document,
        "PowerShell",
        r".\.venv\Scripts\python.exe -m uvicorn app.api:app --host 127.0.0.1 --port 8000",
        "F5F7F6",
    )
    add_text(document, "إعداد CPU الافتراضي:", bold=True, color=GREEN_DARK)
    add_bullet(document, 'WHISPER_MODEL="small"')
    add_bullet(document, 'WHISPER_DEVICE="cpu"')
    add_bullet(document, 'WHISPER_COMPUTE_TYPE="int8"')
    add_text(document, "إعداد GPU المقترح:", bold=True, color=GREEN_DARK)
    add_bullet(document, 'WHISPER_DEVICE="cuda"')
    add_bullet(document, 'WHISPER_COMPUTE_TYPE="float16"')

    add_heading(document, "11", "القيود والتطوير المستقبلي")
    add_bullet(document, "الدقة تتأثر بجودة الميكروفون والضوضاء واللهجة والمصطلحات القانونية.")
    add_bullet(document, "النموذج small يوازن بين السرعة والدقة؛ يمكن استخدام medium أو large-v3 على GPU.")
    add_bullet(document, "أول طلب بعد تشغيل الخادم أبطأ بسبب تحميل النموذج.")
    add_bullet(document, "يمكن تطوير قائمة اختيار الميكروفون بإضافة مؤشر مستوى صوت مرئي واختبار مباشر للجهاز.")
    add_bullet(document, "يمكن إضافة قاموس تصحيح لأسماء القوانين والمواد والأرقام العراقية.")
    add_bullet(document, "يمكن مستقبلاً تنفيذ نص حي streaming، لكنه يزيد تعقيد WebSocket وإدارة الجلسات.")

    add_heading(document, "12", "الخلاصة")
    add_text(
        document,
        "تم دمج خاصية عربية محلية لتحويل الكلام إلى نص في منصة جلسة مع فصل واضح بين التسجيل والتحويل "
        "ونظام الإجابة القانونية. يوفر التنفيذ خصوصية عالية، مراجعة بشرية قبل الإرسال، معالجة أخطاء عملية، "
        "ومساراً قابلاً للتوسع إلى GPU أو نماذج أكبر عند الانتقال إلى بيئة الإنتاج.",
        size=12,
    )
    add_info_box(
        document,
        "الحالة النهائية",
        "الخاصية منفذة ومختبرة وقابلة للتشغيل محلياً، مع توصيات أمنية وتشغيلية واضحة قبل النشر العام.",
    )

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    build_report()
    print("Word report generated successfully.")
