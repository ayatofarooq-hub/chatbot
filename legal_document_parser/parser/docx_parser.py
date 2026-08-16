"""DOCX reader for the independent parser.

This module reads Word documents only. It does not import chatbot, RAG,
embedding, vector-store, database, or model code.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document


@dataclass(frozen=True)
class ParsedDocx:
    source_path: Path
    filename: str
    paragraphs: list[str]
    tables: list[list[list[str]]]
    full_text: str


class DocxParser:
    """Extract paragraphs, tables, and complete text from a DOCX file."""

    def parse(self, source_path: Path) -> ParsedDocx:
        if source_path.suffix.lower() != ".docx":
            raise ValueError(f"Unsupported input file: {source_path.name}")

        document = Document(source_path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        paragraphs = [paragraph for paragraph in paragraphs if paragraph]
        tables = self._extract_tables(document)
        table_text = self._table_text(tables)
        full_text = "\n".join([*paragraphs, *table_text]).strip()

        return ParsedDocx(
            source_path=source_path,
            filename=source_path.name,
            paragraphs=paragraphs,
            tables=tables,
            full_text=full_text,
        )

    def _extract_tables(self, document: Document) -> list[list[list[str]]]:
        tables: list[list[list[str]]] = []
        for table in document.tables:
            rows: list[list[str]] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                cleaned = [cell for cell in cells if cell]
                if cleaned:
                    rows.append(cleaned)
            if rows:
                tables.append(rows)
        return tables

    def _table_text(self, tables: list[list[list[str]]]) -> list[str]:
        lines: list[str] = []
        for table in tables:
            for row in table:
                value = " | ".join(cell for cell in row if cell)
                if value:
                    lines.append(value)
        return lines

