"""Grounded Qwen answers from retrieved legal chunks."""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import ollama
from docx import Document

from app.citation_registry import filter_results_to_registered, load_registry
from app.config import PROJECT_ROOT
from app.runtime_settings import runtime_settings
from app.search_index import search
from app.legal_source_text import NO_USABLE_LEGAL_SOURCE_TEXT, source_text_from_payload
from app.retrieval_logging import (
    log_answer_generation,
    log_long_text_status,
    log_retrieval_results,
)
from app.text_encoding import repair_mojibake
from .json_loader import load_legal_json_documents, primary_document_text


INSUFFICIENT_CONTEXT = "لم أجد نصًا قانونيًا كافيًا للإجابة عن هذا السؤال في المستندات المتاحة."
FULL_TEXT_HEADING = "النص القانوني الكامل:"
FULL_TEXT_FALLBACK_ANSWER = "تم العثور على النص القانوني الآتي في المصدر الأقرب للسؤال."
MISSING_DECISION_NUMBER_ANSWER = "رقم القرار غير مثبت في النص المستخرج من الوثيقة."
RELATED_TOPICS_HEADING = "مواضيع مقترحة من نفس النص:"
JSON_TEXT_DIRS = (
    PROJECT_ROOT / "legal_document_parser" / "output" / "json",
    PROJECT_ROOT / "data" / "output",
)
DOCX_INPUT_DIRS = (
    PROJECT_ROOT / "legal_document_parser" / "input" / "docx",
)

SYSTEM_INSTRUCTION = """
أنت مساعد قانوني عراقي.

أجب اعتمادًا على النصوص القانونية المسترجعة فقط.
لا تخترع أي معلومة غير موجودة في المصادر.

إذا لم تجد الإجابة في النصوص المسترجعة، قل بوضوح:
"لم أجد نصًا قانونيًا كافيًا للإجابة عن هذا السؤال في المستندات المتاحة."

عند وجود نص قانوني مباشر، استخدمه كأساس للإجابة.
لا تغير المعنى القانوني للنص.
لا تستخدم المعرفة العامة خارج النصوص المسترجعة.
إذا سأل المستخدم عن مبلغ، فابحث في النصوص عن كلمة "بمبلغ" أو "مقداره" أو رقم متبوع بعملة، وانقل المبلغ كما ورد.
لا تقل إنك لم تجد الإجابة إذا كان النص المسترجع يحتوي عبارة مباشرة تجيب السؤال.
اكتب جوابًا مباشرًا فقط، ولا تبدأ بعبارات عامة مثل "وفقًا للنصوص المسترجعة".
For Word/docx-derived documents, always answer with a brief summary only.
Do not quote, append, or reproduce the full document text. Keep the answer to
one short paragraph or at most three concise bullets, while preserving exact
numbers, dates, parties, and obligations that directly answer the question.
End useful answers with "مواضيع مقترحة من نفس النص:" followed by up to three
related follow-up questions extracted from facts that appear in the retrieved
Word/docx legal text itself.
Phrase the suggestions interactively, for example "هل تريد معرفة..." or
"أستطيع مساعدتك في...".
""".strip()


@dataclass(frozen=True)
class RetrievedSource:
    document_id: str
    document: str
    section: str
    item: str
    chunk_id: str
    text: str
    paragraph_indexes: tuple[int, ...] = ()
    json_path: str = ""

    def public_dict(self) -> dict[str, str]:
        """Return only the source fields exposed to callers."""

        return {
            "document": self.document,
            "section": self.section,
            "item": self.item,
            "chunk_id": self.chunk_id,
        }


def _clean(value: Any) -> str:
    return repair_mojibake(str(value or "")).strip()


def _strip_insufficient_context(answer: str) -> str:
    insufficient_context_variants = (
        INSUFFICIENT_CONTEXT,
        "\u0644\u0645 \u062a\u062c\u062f \u0646\u0635\u064b\u0627 \u0642\u0627\u0646\u0648\u0646\u064a\u064b\u0627 \u0643\u0627\u0641\u064a\u064b\u0627 \u0644\u0644\u0625\u062c\u0627\u0628\u0629 \u0639\u0646 \u0647\u0630\u0627 \u0627\u0644\u0633\u0624\u0627\u0644 \u0641\u064a \u0627\u0644\u0645\u0633\u062a\u0646\u062f\u0627\u062a \u0627\u0644\u0645\u062a\u0627\u062d\u0629.",
    )
    cleaned = _clean(answer)
    for phrase in insufficient_context_variants:
        cleaned = cleaned.replace(_clean(phrase), "")
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip(" \t\r\n.-")


def asks_for_decision_number(question: str) -> bool:
    """Return whether the question asks specifically for the decision number."""

    normalized = _clean(question)
    return bool(
        re.search(r"(?:ما|ماهو|ما\s+هو|اذكر|اعطني|أعطني).{0,30}رقم\s+القرار", normalized)
        or re.search(r"رقم\s+قرار\s+مجلس\s+الوزراء", normalized)
    )


def _metadata_decision_number(metadata: dict[str, Any]) -> str:
    return _clean(metadata.get("decision_number"))


def extracted_sources(results: dict) -> list[RetrievedSource]:
    """Convert retrieved Chroma results to traceable source records."""

    documents = results.get("documents", [[]])[0]
    metadatas = results.get("metadatas", [[]])[0]
    sources: list[RetrievedSource] = []
    for document, metadata in zip(documents, metadatas):
        text = _clean(document)
        chunk_id = _clean(metadata.get("chunk_id") or metadata.get("id"))
        if not text or not chunk_id:
            continue
        sources.append(
            RetrievedSource(
                document_id=_clean(metadata.get("document_id")),
                document=_clean(metadata.get("source_file")),
                section=_clean(metadata.get("section") or metadata.get("section_title")),
                item=_clean(metadata.get("item_number") or metadata.get("article_reference")),
                chunk_id=chunk_id,
                text=text,
                paragraph_indexes=_parse_paragraph_indexes(metadata.get("paragraph_indexes")),
                json_path=_clean(metadata.get("json_path") or metadata.get("original_json_path")),
            )
        )
    return sources


def _parse_paragraph_indexes(value: Any) -> tuple[int, ...]:
    """Parse paragraph index metadata stored as scalar Chroma metadata."""

    indexes = []
    if isinstance(value, int):
        return (value,)
    for item in re.findall(r"\d+", str(value or "")):
        try:
            indexes.append(int(item))
        except ValueError:
            continue
    return tuple(dict.fromkeys(indexes))


def build_source_context(sources: list[RetrievedSource]) -> str:
    """Build the required SOURCE n context block for Qwen."""

    blocks = []
    for index, source in enumerate(sources, start=1):
        blocks.append(
            "\n".join(
                [
                    f"SOURCE {index}",
                    f"Document: {source.document}",
                    f"Section: {source.section}",
                    f"Item: {source.item}",
                    "Text:",
                    source.text,
                ]
            )
        )
    return "\n\n".join(blocks)


def _metadata_at(results: dict, index: int) -> dict[str, Any]:
    metadatas = results.get("metadatas", [[]])[0]
    if index < len(metadatas) and isinstance(metadatas[index], dict):
        return metadatas[index]
    return {}


def _metadata_text(metadata: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = _clean(metadata.get(key))
        if value:
            return value
    return ""


def _split_metadata_values(*values: Any) -> list[str]:
    items: list[str] = []
    for value in values:
        if isinstance(value, list):
            candidates = value
        else:
            candidates = re.split(r"\s*\|\s*|\n+", str(value or ""))
        for candidate in candidates:
            text = _clean(candidate)
            if text:
                items.append(text)
    return list(dict.fromkeys(items))


def _format_answer_context_list(label: str, values: list[str]) -> list[str]:
    if not values:
        return [f"{label}:", "غير مثبت"]
    return [f"{label}:", *values]


def build_document_answer_context(results: dict, sources: list[RetrievedSource]) -> str:
    """Build a structured document context used as the answer generator source."""

    full_text_records = full_texts_for_sources(sources)
    relevance_scores = results.get("relevance_scores", [[]])[0]
    full_text_by_chunk = {
        str(record.get("chunk_id") or ""): str(record.get("original_long_text") or record.get("full_text") or "")
        for record in full_text_records
    }
    blocks: list[str] = []
    for index, source in enumerate(sources, start=1):
        metadata = _metadata_at(results, index - 1)
        relevance_score = relevance_scores[index - 1] if index <= len(relevance_scores) else ""
        references = _split_metadata_values(
            metadata.get("document_reference_numbers"),
            metadata.get("reference_numbers"),
            metadata.get("recommendation_numbers"),
            metadata.get("legal_reference"),
        )
        entities = _split_metadata_values(metadata.get("entities"), metadata.get("organizations"), metadata.get("companies"))
        full_source_text = full_text_by_chunk.get(source.chunk_id) or "غير مثبت في JSON الأصلي"
        blocks.append(
            "\n".join(
                [
                    f"DOCUMENT {index}",
                    "",
                    "Source:",
                    source.document or "غير مثبت",
                    "",
                    "Relevance Score:",
                    _clean(relevance_score) or "غير مثبت",
                    "",
                    "Title:",
                    _metadata_text(metadata, "title", "document_title") or source.document or "غير مثبت",
                    "",
                    "Type:",
                    _metadata_text(metadata, "document_type", "classification") or "غير مثبت",
                    "",
                    "Year:",
                    _metadata_text(metadata, "year", "law_year") or "غير مثبت",
                    "",
                    "Issue Date:",
                    _metadata_text(metadata, "issue_date") or "غير مثبت",
                    "",
                    "Session:",
                    _metadata_text(metadata, "session_number") or "غير مثبت",
                    "",
                    "Session Date:",
                    _metadata_text(metadata, "session_date") or "غير مثبت",
                    "",
                    *_format_answer_context_list("References", references),
                    "",
                    *_format_answer_context_list("Entities", entities),
                    "",
                    "RELEVANT RETRIEVED TEXT:",
                    source.text,
                    "",
                    "FULL SOURCE TEXT:",
                    full_source_text,
                ]
            )
        )
    context = "\n\n".join(blocks)
    if len(blocks) > 1:
        instruction = (
            "MULTI-DOCUMENT SOURCE RULE:\n"
            "Mention the supporting source for each distinct finding using DOCUMENT number, title, or source label."
        )
        return f"{instruction}\n\n{context}"
    return context


def wants_full_document_context(question: str) -> bool:
    """Return whether the question asks for the whole source document."""

    normalized = _clean(question)
    patterns = (
        r"\b(?:النص|نص)\s+(?:الكامل|كامل[اً]?)\b",
        r"\b(?:اعطني|أعطني|اريد|أريد)\s+(?:نص|النص).{0,30}(?:كامل|كاملا|كاملًا)\b",
        r"\b(?:مضمون|ملخص|خلاصة)\s+(?:هذا\s+)?(?:القرار|الوثيقة|المستند)\b",
        r"\bما\s+(?:هو\s+)?(?:مضمون|محتوى)\s+(?:هذا\s+)?(?:القرار|الوثيقة|المستند)\b",
        r"\bما\s+الذي\s+قرره\s+مجلس\s+الوزراء\b",
        r"\bماذا\s+قرر\s+مجلس\s+الوزراء\b",
    )
    return any(re.search(pattern, normalized) for pattern in patterns)


def wants_exact_full_document_text(question: str) -> bool:
    normalized = _clean(question)
    return bool(
        re.search(r"\b(?:اعطني|أعطني|اريد|أريد)\s+(?:نص|النص).{0,30}(?:كامل|كاملا|كاملًا)\b", normalized)
        or re.search(r"\b(?:النص|نص)\s+(?:الكامل|كامل[اً]?)\b", normalized)
    )


def product_price_answer_from_sources(question: str, sources: list[RetrievedSource]) -> str:
    """Extract product price changes from retrieved source text."""

    question_text = _clean(question)
    asks_previous_price = bool(re.search(r"(?:كان|سابق|سابقًا|سابقا|قبل)", question_text))
    product_match = re.search(r"سعر\s+منتوج\s+([\u0600-\u06FF\s]+?)(?:\?|؟|$)", question_text)
    if not product_match:
        product_match = re.search(r"منتوج\s+([\u0600-\u06FF\s]+?)(?:\?|؟|$)", question_text)
    if not product_match:
        product_match = re.search(r"سعر\s+([\u0600-\u06FF\s]+?)(?:\?|؟|$)", question_text)
    if not product_match:
        product_match = re.search(r"بشأن\s+([\u0600-\u06FF\s]+?)(?:\?|؟|$)", question_text)
    product = re.sub(r"\s+", " ", product_match.group(1)).strip() if product_match else ""
    product = re.sub(r"\b(?:سابقًا|سابقا|سابق|الآن|حاليًا|حاليا)\b", "", product).strip()
    if not product:
        return ""

    product_terms = {term for term in re.findall(r"[\u0600-\u06FF]+", product) if len(term) > 2}
    price_pattern = re.compile(
        r"منتوج\s+([\u0600-\u06FF\s]+?)\s*\(([^)]+)\)\s*دينار\s*/\s*([^\s.،]+)\s*"
        r"بدل\S*\s+من\s*\(([^)]+)\)\s*دينار\s*/\s*([^\s.،]+)",
    )
    for source in sources:
        for match in price_pattern.finditer(source.text):
            matched_product, new_price, new_unit, old_price, old_unit = [_clean(value) for value in match.groups()]
            matched_terms = set(re.findall(r"[\u0600-\u06FF]+", matched_product))
            if product_terms and not product_terms <= matched_terms:
                continue
            if asks_previous_price:
                return f"كان سعر منتوج {matched_product} سابقًا {old_price} دينار / {old_unit}."
            return (
                f"تم تعديل سعر منتوج {matched_product} ليصبح {new_price} دينار / {new_unit} "
                f"بدلًا من {old_price} دينار / {old_unit}."
            )
    return ""


def sources_with_full_text_context(sources: list[RetrievedSource]) -> list[RetrievedSource]:
    """Replace the best source text with original long_text for full-document questions."""

    if not sources:
        return []
    full_text_records = full_texts_for_sources(sources[:1])
    by_chunk_id = {
        str(record.get("chunk_id") or ""): str(record.get("original_long_text") or record.get("full_text") or "")
        for record in full_text_records
    }
    expanded = []
    for source in sources:
        full_text = by_chunk_id.get(source.chunk_id)
        if full_text:
            expanded.append(
                RetrievedSource(
                    document_id=source.document_id,
                    document=source.document,
                    section=source.section,
                    item=source.item,
                    chunk_id=source.chunk_id,
                    text=full_text,
                    paragraph_indexes=source.paragraph_indexes,
                    json_path=source.json_path,
                )
            )
        else:
            expanded.append(source)
    return expanded


def build_user_message(question: str, context: str) -> str:
    """Build the user message sent to Qwen."""

    return "\n\n".join(
        [
            "السؤال:",
            question,
            (
                "For Word/docx sources, provide a summary only. Do not include or "
                "append the full legal text, even when full source text appears in "
                "the retrieved context."
            ),
            (
                'End the answer with "مواضيع مقترحة من نفس النص:" and up to '
                "three follow-up questions extracted from facts in the source text."
            ),
            "النصوص القانونية المسترجعة:",
            context,
            (
                "أجب من النصوص أعلاه فقط. إذا كان الجواب موجودًا، استخرج الجواب "
                "مباشرة من النص القانوني دون إضافة معلومات خارجية. إذا كان السؤال "
                "عن مبلغ، فانقل الرقم والعملة والكتابة اللفظية للمبلغ كما وردت في النص."
            ),
        ]
    )


def call_qwen(question: str, context: str, status_callback: Callable[[str], None] | None = None) -> str:
    """Call the existing configured Qwen/Ollama chat model."""

    settings = runtime_settings()["model"]
    client = ollama.Client(
        host=settings["ollama_base_url"],
        timeout=settings["request_timeout"],
    )
    if status_callback:
        status_callback("جارٍ توليد إجابة قانونية من النصوص المسترجعة...")
    response = client.chat(
        model=settings["chat_model"],
        messages=[
            {"role": "system", "content": SYSTEM_INSTRUCTION},
            {"role": "user", "content": build_user_message(question, context)},
        ],
        stream=False,
        keep_alive=settings["keep_alive"],
        options={
            "temperature": 0.0,
            "top_p": settings["top_p"],
            "num_ctx": settings["context_length"],
            "num_predict": settings["max_answer_tokens"],
        },
    )
    return _clean(response["message"]["content"])


def estimate_confidence(answer: str, sources: list[RetrievedSource], results: dict) -> float:
    """Return a conservative confidence score for grounded answers."""

    if not sources or not answer or INSUFFICIENT_CONTEXT in answer:
        return 0.0
    relevance_scores = results.get("relevance_scores", [[]])[0]
    if relevance_scores:
        try:
            return round(max(0.0, min(1.0, float(relevance_scores[0]))), 2)
        except (TypeError, ValueError):
            pass
    return 0.7


def _answer_numbers(answer: str) -> list[str]:
    return re.findall(r"\d[\d.,/]*", answer)


def supporting_sources(answer: str, sources: list[RetrievedSource]) -> list[RetrievedSource]:
    """Return only retrieved sources that directly support the answer."""

    if not sources or not answer:
        return []
    if INSUFFICIENT_CONTEXT in answer:
        return sources[:1]

    numbers = _answer_numbers(answer)
    if numbers:
        matching = [
            source
            for source in sources
            if any(number in source.text for number in numbers)
        ]
        if matching:
            return matching

    answer_terms = {
        term
        for term in re.findall(r"[\w\u0600-\u06ff]+", answer)
        if len(term) > 2
    }
    matching = []
    for source in sources:
        source_terms = set(re.findall(r"[\w\u0600-\u06ff]+", source.text))
        if answer_terms and len(answer_terms & source_terms) >= min(3, len(answer_terms)):
            matching.append(source)
    return matching or sources[:1]


def _source_keys(*values: Any) -> set[str]:
    keys = set()
    for value in values:
        text = _clean(value)
        if not text:
            continue
        keys.add(text)
        keys.add(Path(text).name)
        keys.add(Path(text).stem)
    return keys


def _texts_from_items(items: Any) -> list[str]:
    texts = []
    if not isinstance(items, list):
        return texts
    for item in items:
        if isinstance(item, dict):
            text = _clean(item.get("text") or item.get("body") or item.get("content"))
        else:
            text = _clean(item)
        if text:
            texts.append(text)
    return texts


def _full_text_from_payload(payload: Any) -> str:
    if not isinstance(payload, dict):
        return ""
    text = _clean(source_text_from_payload(payload))
    if text:
        return text

    if payload.get("schema_version") == "iraqi_legal_document.v2":
        return ""

    for key in ("full_text", "text", "content"):
        legacy_text = _clean(payload.get(key))
        if legacy_text:
            return legacy_text

    parts = []
    parts.extend(_texts_from_items(payload.get("paragraphs")))
    parts.extend(_texts_from_items(payload.get("sections")))

    decision = payload.get("decision")
    if isinstance(decision, dict):
        parts.extend(_texts_from_items(decision.get("sections")))
        parts.extend(_texts_from_items(decision.get("numbered_items")))

    return "\n".join(dict.fromkeys(part for part in parts if part)).strip()


def _full_text_from_docx(path: Path) -> str:
    """Extract complete text directly from an original input DOCX file."""

    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def load_full_text_index() -> dict[str, str]:
    """Index complete source text from original JSON records and Word files."""

    index: dict[str, str] = {}

    for document in load_legal_json_documents():
        text = document.long_text.strip()
        for key in _source_keys(document.document_id, document.source_file, document.source_path.name, str(document.source_path)):
            index[key] = text

    for folder in JSON_TEXT_DIRS:
        if not folder.exists():
            continue
        for path in sorted(folder.glob("*.json")):
            try:
                payload = json.loads(path.read_text(encoding="utf-8-sig"))
            except (OSError, json.JSONDecodeError):
                continue
            text = _full_text_from_payload(payload)
            if not text:
                continue
            source = payload.get("source") if isinstance(payload.get("source"), dict) else {}
            metadata = payload.get("metadata") if isinstance(payload.get("metadata"), dict) else {}
            document = payload.get("document") if isinstance(payload.get("document"), dict) else {}
            for key in _source_keys(
                path.name,
                path.stem,
                payload.get("source_file"),
                payload.get("document_id"),
                source.get("filename"),
                source.get("source_file"),
                metadata.get("source_filename"),
                metadata.get("document_id"),
                document.get("id"),
            ):
                index.setdefault(key, text)

    for folder in DOCX_INPUT_DIRS:
        if not folder.exists():
            continue
        for path in sorted(folder.glob("*.docx")):
            if path.name.startswith("~$"):
                continue
            try:
                text = _clean(_full_text_from_docx(path))
            except Exception:
                continue
            if not text:
                continue
            for key in _source_keys(path.name, path.stem, str(path)):
                index.setdefault(key, text)

    return index


def full_texts_for_sources(sources: list[RetrievedSource]) -> list[dict[str, Any]]:
    """Return complete source text records plus relevant paragraph linkage."""

    if not sources:
        return []
    full_text_index = load_full_text_index()

    records = []
    seen = set()
    for source in sources:
        text = ""
        for key in _source_keys(source.document_id, source.document, source.json_path):
            text = full_text_index.get(key, "")
            if text:
                break
        if not text or text in seen:
            continue
        seen.add(text)
        records.append(
            {
                **source.public_dict(),
                "full_text": text,
                "full_text_length": len(text),
                "original_long_text": text,
                "original_long_text_length": len(text),
                "relevant_text": source.text,
                "paragraph_indexes": list(source.paragraph_indexes),
                "document_id": source.document_id,
                "filename": source.document,
            }
        )
    return records


def _answerable_tokens(text: str) -> set[str]:
    """Return meaningful tokens for direct full-text fallback matching."""

    tokens = set()
    for token in re.findall(r"[\w\u0600-\u06ff]+", _clean(text)):
        if len(token) <= 2:
            continue
        if token in {"ما", "من", "عن", "على", "في", "هل", "الى", "إلى", "اريد", "أريد"}:
            continue
        tokens.add(token)
    return tokens


def _full_text_match_score(question: str, text: str, source_key: str) -> float:
    question_tokens = _answerable_tokens(question)
    if not question_tokens:
        return 0.0
    haystack_tokens = _answerable_tokens(f"{source_key}\n{text}")
    if not haystack_tokens:
        return 0.0
    overlap = len(question_tokens & haystack_tokens) / len(question_tokens)
    number_bonus = 0.0
    for number in re.findall(r"\d[\d/.,-]*", question):
        if number and number in text:
            number_bonus += 0.25
    return min(1.0, overlap + number_bonus)


def fallback_results_from_full_text(question: str, limit: int = 3) -> dict:
    """Build retrieval-like results by searching complete Word/JSON source text."""

    scored: list[tuple[float, str, str]] = []
    seen_texts = set()
    for key, text in load_full_text_index().items():
        if not text or text in seen_texts:
            continue
        seen_texts.add(text)
        score = _full_text_match_score(question, text, key)
        if score <= 0:
            continue
        scored.append((score, key, text))

    scored.sort(key=lambda item: item[0], reverse=True)
    documents = []
    metadatas = []
    distances = []
    relevance_scores = []
    for score, key, text in scored[:limit]:
        source_name = Path(key).name or key
        chunk_id = f"full_text_fallback:{source_name}"
        documents.append(text)
        metadatas.append(
            {
                "chunk_id": chunk_id,
                "document_id": Path(source_name).stem,
                "source_file": source_name,
                "filename": source_name,
                "retrieval_layer": "full_text_fallback",
                "full_source_resolver": "full_text",
            }
        )
        distances.append(max(0.0, 1.0 - score))
        relevance_scores.append(score)

    return {
        "documents": [documents],
        "metadatas": [metadatas],
        "distances": [distances],
        "relevance_scores": [relevance_scores],
    }


def extractive_summary_from_sources(question: str, sources: list[RetrievedSource]) -> str:
    """Return a brief source-grounded summary when the model declines answerable text."""

    question_tokens = _answerable_tokens(question)
    candidates: list[tuple[float, str]] = []
    for source in sources:
        for raw_line in re.split(r"[\n\r]+", source.text):
            line = re.sub(r"\s+", " ", _clean(raw_line)).strip()
            if len(line) < 12:
                continue
            line_tokens = _answerable_tokens(line)
            score = len(question_tokens & line_tokens) / len(question_tokens) if question_tokens else 0.0
            if any(number in line for number in re.findall(r"\d[\d/.,-]*", question)):
                score += 0.25
            candidates.append((score, line))

    candidates.sort(key=lambda item: item[0], reverse=True)
    selected = []
    for score, line in candidates:
        if score <= 0 and selected:
            continue
        if line in selected:
            continue
        selected.append(line)
        if len(selected) == 3:
            break

    if not selected:
        return ""
    if len(selected) == 1:
        return selected[0]
    return "\n".join(f"- {line}" for line in selected)


def answer_decision_number_if_known(question: str, results: dict) -> dict[str, Any] | None:
    """Answer decision-number questions without inferring from reference numbers."""

    if not asks_for_decision_number(question):
        return None
    sources = extracted_sources(results)
    metadatas = results.get("metadatas", [[]])[0]
    metadata = metadatas[0] if metadatas and isinstance(metadatas[0], dict) else {}
    decision_number = _metadata_decision_number(metadata)
    public_sources = [source.public_dict() for source in sources[:1]]
    full_text_sources = full_texts_for_sources(sources[:1])
    full_text = "\n\n".join(record["full_text"] for record in full_text_sources)
    if not decision_number:
        return {
            "answer": MISSING_DECISION_NUMBER_ANSWER,
            "sources": public_sources,
            "full_text": full_text,
            "full_text_sources": full_text_sources,
            "confidence": 1.0,
        }
    return {
        "answer": f"رقم القرار هو {decision_number}.",
        "sources": public_sources,
        "full_text": full_text,
        "full_text_sources": full_text_sources,
        "confidence": 1.0,
    }


def _top_metadata(results: dict) -> dict[str, Any]:
    metadatas = results.get("metadatas", [[]])[0]
    return metadatas[0] if metadatas and isinstance(metadatas[0], dict) else {}


def _metadata_requires_original_source_text(metadata: dict[str, Any]) -> bool:
    if not metadata:
        return False
    if metadata.get("json_path") or metadata.get("original_json_path"):
        return True
    if str(metadata.get("source_type") or "") == "legal_document_parser_json":
        return True
    if str(metadata.get("full_source_resolver") or "") == "json":
        return True
    return False


def answer_unusable_source_text_if_needed(
    results: dict,
    sources: list[RetrievedSource],
    full_text_sources: list[dict[str, Any]],
) -> dict[str, Any] | None:
    """Return a clear error when a JSON source lacks long_text/body/paragraphs."""

    metadata = _top_metadata(results)
    if not _metadata_requires_original_source_text(metadata):
        return None
    if full_text_sources and full_text_sources[0].get("full_text"):
        return None
    cited_sources = sources[:1]
    return {
        "answer": NO_USABLE_LEGAL_SOURCE_TEXT,
        "sources": [source.public_dict() for source in cited_sources],
        "full_text": "",
        "full_text_sources": [],
        "confidence": 0.0,
    }


def requested_document_date_field(question: str) -> tuple[str, str] | None:
    """Return the precise date metadata field requested by the question."""

    text = _clean(question)
    if re.search(r"(?:متى|تاريخ).{0,30}(?:عقدت|انعقدت|انعقاد).{0,20}الجلسة", text):
        return "session_date", "تاريخ انعقاد الجلسة"
    if re.search(r"(?:متى|تاريخ).{0,30}(?:صدر|صدور|اصدار|إصدار).{0,20}(?:القرار|الوثيقة)", text):
        return "issue_date", "تاريخ صدور القرار"
    return None


def answer_document_date_if_known(
    question: str,
    results: dict,
    sources: list[RetrievedSource],
) -> dict[str, Any] | None:
    """Answer issue/session date questions without confusing the two fields."""

    requested = requested_document_date_field(question)
    if not requested:
        return None
    field, label = requested
    value = _clean(_top_metadata(results).get(field))
    cited_sources = sources[:1]
    full_text_sources = full_texts_for_sources(cited_sources)
    full_text = "\n\n".join(record["full_text"] for record in full_text_sources)
    if value:
        answer = f"{label}: {value}."
    else:
        answer = f"{label} غير مثبت في النص المستخرج من الوثيقة."
    return {
        "answer": answer,
        "sources": [source.public_dict() for source in cited_sources],
        "full_text": full_text,
        "full_text_sources": full_text_sources,
        "confidence": 1.0 if value else 0.7,
    }


def asks_for_recommendation_number(question: str) -> bool:
    text = _clean(question)
    return bool(re.search(r"رقم.{0,30}توصية.{0,40}المجلس\s+الوزاري\s+للاقتصاد", text))


def _recommendation_number_from_text(text: str) -> str:
    match = re.search(r"توصية\s+المجلس\s+الوزاري\s+للاقتصاد\s*\(\s*([^)]+?)\s*\)", _clean(text))
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def answer_recommendation_number_if_known(
    question: str,
    results: dict,
    sources: list[RetrievedSource],
) -> dict[str, Any] | None:
    """Answer recommendation-number questions without using decision_number."""

    if not asks_for_recommendation_number(question):
        return None
    cited_sources = sources[:1]
    full_text_sources = full_texts_for_sources(cited_sources)
    metadata_value = _clean(_top_metadata(results).get("recommendation_numbers"))
    if not metadata_value:
        for record in full_text_sources:
            metadata_value = _recommendation_number_from_text(
                str(record.get("original_long_text") or record.get("full_text") or "")
            )
            if metadata_value:
                break
    if not metadata_value:
        for source in cited_sources:
            metadata_value = _recommendation_number_from_text(source.text)
            if metadata_value:
                break
    answer = (
        f"رقم توصية المجلس الوزاري للاقتصاد: {metadata_value}."
        if metadata_value
        else "رقم توصية المجلس الوزاري للاقتصاد غير مثبت في النص المستخرج من الوثيقة."
    )
    return {
        "answer": answer,
        "sources": [source.public_dict() for source in cited_sources],
        "full_text": "\n\n".join(record["full_text"] for record in full_text_sources),
        "full_text_sources": full_text_sources,
        "confidence": 1.0 if metadata_value else 0.7,
    }


def answer_full_document_text_if_requested(
    question: str,
    sources: list[RetrievedSource],
) -> dict[str, Any] | None:
    """Do not return Word/docx long_text directly; the answer must be a summary."""

    return None


def asks_for_source_book(question: str) -> bool:
    """Return whether the question asks which official book a decision relied on."""

    text = _clean(question)
    return bool(
        re.search(r"(?:بناء|استناد).{0,20}(?:على|إلى|الى).{0,40}(?:اي|أي)\s+كتاب", text)
        or re.search(r"(?:اي|أي)\s+كتاب.{0,60}(?:تعديل|عدلت|تعدلت)", text)
        or re.search(r"كتاب.{0,40}(?:تم|جرى).{0,40}تعديل", text)
        or re.search(r"(?:رقم|تاريخ).{0,20}كتاب", text)
    )


def _reference_book_from_text(text: str) -> str:
    normalized_text = _clean(text)
    pattern = re.compile(
        r"كتاب\s+(.{0,80}?)\s+المرقم\s+بالعدد\s*\(\s*([^)]+?)\s*\)\s+المؤرخ\s+في\s+([0-9٠-٩/\\-]+)",
        flags=re.DOTALL,
    )
    for match in pattern.finditer(normalized_text):
        issuer, number, date = [re.sub(r"\s+", " ", value).strip(" .،؛:") for value in match.groups()]
        if not number or not date:
            continue
        issuer_text = f"كتاب {issuer}" if issuer else "الكتاب"
        return f"{issuer_text} المرقم بالعدد ({number}) المؤرخ في {date}."
    return ""


def _reference_book_parts_from_text(text: str) -> tuple[str, str, str] | None:
    normalized_text = _clean(text)
    pattern = re.compile(
        r"كتاب\s+(.{0,80}?)\s+المرقم\s+بالعدد\s*\(\s*([^)]+?)\s*\)\s+المؤرخ\s+في\s+([0-9٠-٩/\\-]+)",
        flags=re.DOTALL,
    )
    match = pattern.search(normalized_text)
    if not match:
        return None
    issuer, number, date = [re.sub(r"\s+", " ", value).strip(" .،؛:") for value in match.groups()]
    return issuer, number, date


def answer_source_book_if_known(
    question: str,
    sources: list[RetrievedSource],
) -> dict[str, Any] | None:
    """Answer source-book questions from original full text when available."""

    if not asks_for_source_book(question):
        return None
    cited_sources = sources[:1]
    full_text_sources = full_texts_for_sources(cited_sources)
    lookup_texts = [
        str(record.get("original_long_text") or record.get("full_text") or "")
        for record in full_text_sources
    ]
    lookup_texts.extend(source.text for source in cited_sources)
    question_text = _clean(question)
    for text in lookup_texts:
        parts = _reference_book_parts_from_text(text)
        if parts:
            issuer, number, date = parts
            issuer_label = issuer or "الكتاب"
            if re.search(r"رقم.{0,20}كتاب", question_text):
                answer = f"رقم كتاب {issuer_label}: {number}."
            elif re.search(r"تاريخ.{0,20}كتاب", question_text):
                answer = f"تاريخ كتاب {issuer_label}: {date}."
            else:
                issuer_text = f"كتاب {issuer_label}" if issuer else "الكتاب"
                answer = f"{issuer_text} المرقم بالعدد ({number}) المؤرخ في {date}."
            return {
                "answer": answer,
                "sources": [source.public_dict() for source in cited_sources],
                "full_text": "\n\n".join(record["full_text"] for record in full_text_sources),
                "full_text_sources": full_text_sources,
                "confidence": 1.0,
            }
    return {
        "answer": "الكتاب الذي استند إليه تعديل الأسعار غير مثبت في النص المستخرج من الوثيقة.",
        "sources": [source.public_dict() for source in cited_sources],
        "full_text": "\n\n".join(record["full_text"] for record in full_text_sources),
        "full_text_sources": full_text_sources,
        "confidence": 0.7,
    }


def answer_with_full_legal_text(answer: str, sources: list[RetrievedSource]) -> str:
    """Keep the public answer concise; full text remains in full_text_sources."""

    return answer


def related_topics_for_sources(question: str, sources: list[RetrievedSource]) -> list[str]:
    """Build source-aware follow-up topics without adding legal facts."""

    text = _clean(" ".join([question, *(source.document for source in sources), *(source.text for source in sources)]))
    suggestions: list[str] = []

    def add(topic: str) -> None:
        if topic not in suggestions:
            suggestions.append(topic)

    if "قرار" in text or "مجلس الوزراء" in text:
        add("ما رقم القرار وتاريخ صدوره؟")
        add("ما الجهات أو الأطراف المشمولة بالقرار؟")
        add("ما الالتزامات أو الإجراءات المطلوبة لتنفيذه؟")
    if "قانون" in text:
        add("ما نطاق تطبيق هذا القانون؟")
        add("ما أهم الحقوق أو الالتزامات الواردة فيه؟")
        add("هل توجد عقوبات أو استثناءات مرتبطة بالموضوع؟")
    if "كتاب" in text:
        add("ما رقم وتاريخ الكتاب الذي استندت إليه الوثيقة؟")
    if re.search(r"مبلغ|دينار|دولار|سعر|كلفة|تكلفة", text):
        add("ما المبلغ أو السعر المذكور وبأي عملة؟")

    if not suggestions:
        suggestions = [
            "ما أهم النقاط العملية في هذه الوثيقة؟",
            "ما الجهة المسؤولة أو المعنية بالموضوع؟",
            "ما المعلومات الناقصة التي تحتاج إلى تحقق من المصدر؟",
        ]
    return suggestions[:3]


def _trim_topic_value(value: str, limit: int = 70) -> str:
    value = re.sub(r"\s+", " ", _clean(value)).strip(" .،؛:")
    return value[:limit].rstrip(" .،؛:")


def _source_entity_candidates(text: str) -> list[str]:
    pattern = re.compile(
        r"\b(?:وزارة|شركة|مجلس|هيئة|دائرة|محافظة|لجنة)\s+[\u0600-\u06FF\s]{2,55}",
    )
    entities = []
    for match in pattern.finditer(text):
        entity = _trim_topic_value(re.split(r"[،؛:.\n\r]", match.group(0), 1)[0])
        if entity and entity not in entities:
            entities.append(entity)
    return entities


def related_topics_for_sources(question: str, sources: list[RetrievedSource]) -> list[str]:
    """Build follow-up topics from facts that appear in the same retrieved text."""

    text = _clean(" ".join([*(source.document for source in sources), *(source.text for source in sources)]))
    suggestions: list[str] = []

    def add(topic: str) -> None:
        if topic not in suggestions:
            suggestions.append(topic)

    book = re.search(
        r"كتاب\s+(.{0,80}?)\s+المرقم\s+بالعدد\s*\(\s*([^)]+?)\s*\)\s+المؤرخ\s+في\s+([0-9٠-٩/\\-]+)",
        text,
        flags=re.DOTALL,
    )
    if book:
        issuer, number, date = [_trim_topic_value(value) for value in book.groups()]
        issuer_text = f" {issuer}" if issuer else ""
        add(f"هل تريد معرفة أثر كتاب{issuer_text} المرقم ({number}) المؤرخ في {date}؟")

    amount = re.search(r"([0-9٠-٩][0-9٠-٩.,/ ]+)\s*(دينار|دولار)", text)
    if amount:
        value, currency = [_trim_topic_value(value) for value in amount.groups()]
        add(f"أستطيع مساعدتك في توضيح تفاصيل المبلغ {value} {currency} الوارد في النص.")

    date = re.search(r"\b([0-9٠-٩]{1,2}\s*/\s*[0-9٠-٩]{1,2}\s*/\s*[0-9٠-٩]{4})\b", text)
    if date:
        add(f"هل تريد معرفة دلالة تاريخ {date.group(1)} في هذه الوثيقة؟")

    clause = re.search(r"\b(أولاً|أولًا|ثانياً|ثانيًا|ثالثاً|ثالثًا)\s*[:：]?\s*(.{10,90})", text)
    if clause:
        label = _trim_topic_value(clause.group(1))
        add(f"هل تريد أن أشرح لك البند {label} الوارد في النص؟")

    for entity in _source_entity_candidates(text)[:2]:
        add(f"أستطيع مساعدتك في توضيح دور {entity} في النص.")

    if not suggestions:
        suggestions = [
            "هل تريد معرفة النقطة القانونية الرئيسية التي يقررها هذا النص؟",
            "أستطيع مساعدتك في تحديد العبارة الأهم المرتبطة بسؤالك من نفس النص.",
            "هل تريد أن أراجع لك الجزء الذي يحتاج إلى قراءة تفصيلية من نفس الوثيقة؟",
        ]
    return suggestions[:3]


def append_related_topics(answer: str, question: str, sources: list[RetrievedSource]) -> str:
    """Append interactive follow-up topics to valid grounded answers."""

    if (
        not answer
        or INSUFFICIENT_CONTEXT in answer
        or RELATED_TOPICS_HEADING in answer
        or "مواضيع مقترحة:" in answer
    ):
        return answer
    topics = related_topics_for_sources(question, sources)
    if not topics:
        return answer
    topic_lines = "\n".join(f"- {topic}" for topic in topics)
    return f"{answer.strip()}\n\n{RELATED_TOPICS_HEADING}\n{topic_lines}"


def with_related_topics(result: dict[str, Any], question: str, sources: list[RetrievedSource]) -> dict[str, Any]:
    """Return an answer payload with related topics appended when appropriate."""

    answer = str(result.get("answer") or "")
    return {**result, "answer": append_related_topics(answer, question, sources)}


def answer_from_results(
    question: str,
    results: dict,
    qwen_caller: Callable[[str, str], str] | None = None,
) -> dict[str, Any]:
    """Generate a grounded answer from already retrieved results."""

    registry = load_registry()
    filtered_results, _warnings = filter_results_to_registered(results, registry=registry)
    log_retrieval_results(question, filtered_results)
    sources = extracted_sources(filtered_results)
    if not sources:
        unregistered_sources = extracted_sources(results)
        if unregistered_sources:
            filtered_results = results
            sources = unregistered_sources
            log_answer_generation("using unregistered retrieved Word/doc source")
    if not sources:
        full_text_results = fallback_results_from_full_text(question)
        full_text_sources = extracted_sources(full_text_results)
        if full_text_sources:
            filtered_results = full_text_results
            sources = full_text_sources
            log_retrieval_results(question, filtered_results)
            log_answer_generation("using full-text Word/doc fallback")
    if not sources:
        log_answer_generation("no sources found")
        return {"answer": INSUFFICIENT_CONTEXT, "sources": [], "confidence": 0.0}
    top_metadata = _top_metadata(filtered_results)
    top_full_text_sources = full_texts_for_sources(sources[:1])
    top_full_text_found = bool(top_full_text_sources and top_full_text_sources[0].get("full_text"))
    top_context_length = len(str(top_full_text_sources[0].get("full_text") or "")) if top_full_text_sources else 0
    log_long_text_status(
        top_metadata,
        top_full_text_found,
        "Full document available" if top_full_text_found else "Full document not loaded",
        top_context_length,
    )
    unusable_source_text_answer = answer_unusable_source_text_if_needed(
        filtered_results,
        sources,
        top_full_text_sources,
    )
    if unusable_source_text_answer:
        log_answer_generation(NO_USABLE_LEGAL_SOURCE_TEXT)
        return unusable_source_text_answer
    decision_number_answer = answer_decision_number_if_known(question, filtered_results)
    if decision_number_answer:
        log_answer_generation("direct decision-number answer")
        return with_related_topics(decision_number_answer, question, sources[:1])
    document_date_answer = answer_document_date_if_known(question, filtered_results, sources)
    if document_date_answer:
        log_answer_generation("direct document-date answer")
        return with_related_topics(document_date_answer, question, sources[:1])
    recommendation_number_answer = answer_recommendation_number_if_known(question, filtered_results, sources)
    if recommendation_number_answer:
        log_answer_generation("direct recommendation-number answer")
        return with_related_topics(recommendation_number_answer, question, sources[:1])
    source_book_answer = answer_source_book_if_known(question, sources)
    if source_book_answer:
        log_answer_generation("direct source-book answer")
        return with_related_topics(source_book_answer, question, sources[:1])
    full_document_text_answer = answer_full_document_text_if_requested(question, sources)
    if full_document_text_answer:
        log_answer_generation("direct full-document text answer")
        return full_document_text_answer

    full_document_context = wants_full_document_context(question)
    context_sources = sources_with_full_text_context(sources) if full_document_context else sources
    product_price_answer = product_price_answer_from_sources(question, context_sources)
    if product_price_answer:
        log_answer_generation("direct product-price answer")
        cited_sources = sources[:1] if full_document_context else supporting_sources(product_price_answer, sources)
        full_text_sources = full_texts_for_sources(cited_sources)
        return with_related_topics({
            "answer": product_price_answer,
            "sources": [source.public_dict() for source in cited_sources],
            "full_text": "\n\n".join(record["full_text"] for record in full_text_sources),
            "full_text_sources": full_text_sources,
            "confidence": estimate_confidence(product_price_answer, cited_sources, filtered_results),
        }, question, cited_sources)
    context = build_document_answer_context(filtered_results, context_sources)
    log_long_text_status(
        top_metadata,
        top_full_text_found,
        "Full document loaded" if full_document_context else "Structured context built",
        len(context),
    )
    caller = qwen_caller or (lambda q, c: call_qwen(q, c))
    log_answer_generation("model generation started")
    answer = _clean(caller(question, context))
    if not answer:
        answer = INSUFFICIENT_CONTEXT
    if INSUFFICIENT_CONTEXT in answer and sources:
        fallback_answer = extractive_summary_from_sources(question, context_sources)
        if fallback_answer:
            answer = fallback_answer
    log_answer_generation("model generation completed")
    cited_sources = sources[:1] if full_document_context else supporting_sources(answer, sources)
    full_text_sources = full_texts_for_sources(cited_sources)

    return with_related_topics({
        "answer": answer_with_full_legal_text(answer, cited_sources),
        "sources": [source.public_dict() for source in cited_sources],
        "full_text": "\n\n".join(record["full_text"] for record in full_text_sources),
        "full_text_sources": full_text_sources,
        "confidence": estimate_confidence(answer, cited_sources, filtered_results),
    }, question, cited_sources)


def answer_question(question: str) -> dict[str, Any]:
    """Run hybrid retrieval and answer with Qwen using only retrieved chunks."""

    return answer_from_results(question, search(question))


def main() -> int:
    parser = argparse.ArgumentParser(description="Answer using retrieved legal chunks and Qwen.")
    parser.add_argument("question", nargs="+")
    args = parser.parse_args()
    print(json.dumps(answer_question(" ".join(args.question)), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
