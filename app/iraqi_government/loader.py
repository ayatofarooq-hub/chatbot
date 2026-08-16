"""Document loading for the Iraqi government preprocessing module."""

from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path
import re

from app.iraqi_government.models import LoadedGovernmentDocument


SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".rtf", ".docx", ".pdf"}


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)

    def get_text(self) -> str:
        return "\n".join(self.parts)


def load_document(path: str | Path) -> LoadedGovernmentDocument:
    """Load one supported file into normalized text without touching RAG state."""

    source_path = Path(path)
    if not source_path.exists():
        raise FileNotFoundError(source_path)

    extension = source_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document extension: {extension}")

    if extension in {".txt", ".md"}:
        text = source_path.read_text(encoding="utf-8")
        metadata = {}
    elif extension in {".html", ".htm"}:
        text = _load_html(source_path)
        metadata = {}
    elif extension == ".rtf":
        text = _load_rtf(source_path)
        metadata = {}
    elif extension == ".docx":
        text, metadata = _load_docx(source_path)
    elif extension == ".pdf":
        text, metadata = _load_pdf(source_path)
    else:
        raise ValueError(f"Unsupported document extension: {extension}")

    return LoadedGovernmentDocument(
        source_path=str(source_path),
        source_file=source_path.name,
        filename=source_path.name,
        extension=extension,
        text=normalize_text(text),
        loader_metadata=metadata,
    )


def normalize_text(text: str) -> str:
    """Normalize whitespace while preserving paragraph boundaries."""

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    compacted: list[str] = []
    previous_blank = False
    for line in lines:
        is_blank = not line
        if is_blank and previous_blank:
            continue
        compacted.append(line)
        previous_blank = is_blank
    return "\n".join(compacted).strip()


def _load_html(path: Path) -> str:
    parser = _TextHTMLParser()
    parser.feed(path.read_text(encoding="utf-8"))
    return parser.get_text()


def _load_rtf(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"{\\\*[^{}]*}", " ", raw)
    raw = raw.replace("\\par", "\n")
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    return raw


def _load_docx(path: Path) -> tuple[str, dict[str, str]]:
    from docx import Document

    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        if rows:
            tables.append(rows)
    metadata = {"table_count": str(len(tables))}
    return "\n\n".join(paragraphs), metadata


def _load_pdf(path: Path) -> tuple[str, dict[str, str]]:
    import fitz

    parts: list[str] = []
    with fitz.open(path) as pdf:
        for page in pdf:
            page_text = page.get_text("text").strip()
            if page_text:
                parts.append(page_text)
        metadata = {"page_count": str(pdf.page_count)}
    return "\n\n".join(parts), metadata
