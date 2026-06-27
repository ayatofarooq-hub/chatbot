"""Structured source document loaders for the legal knowledge pipeline."""

from __future__ import annotations

import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

try:
    from .document_classifier import classify_document
    from .text_cleaning import clean_text
except ImportError:
    from document_classifier import classify_document
    from text_cleaning import clean_text


PAGE_SEPARATOR_PATTERN = re.compile(r"^--- PAGE (\d+) ---\s*$", re.MULTILINE)
HEADING_STYLE_PATTERN = re.compile(r"^(?:heading|Ø¹Ù†ÙˆØ§Ù†)\s*(\d+)?", re.IGNORECASE)
ARTICLE_PATTERN = re.compile(
    r"^\s*(?P<label>Ø§Ù„Ù…Ø§Ø¯Ø©|Ù…Ø§Ø¯Ø©)\s*"
    r"(?P<number>[\(\[]?[0-9\u0660-\u0669]+[\)\]]?)",
)
LEGAL_SECTION_PATTERN = re.compile(
    r"^\s*(?P<label>Ø§Ù„Ø¨Ø§Ø¨|Ø§Ù„ÙØµÙ„|Ø§Ù„Ù‚Ø³Ù…|Ø§Ù„ÙØ±Ø¹)\s+"
    r"(?P<name>[^\n:]{1,120})",
)


@dataclass(frozen=True)
class DocumentBlock:
    """One ordered semantic block extracted from a source document."""

    text: str
    block_type: str = "paragraph"
    page_number: int = 1
    heading_level: int | None = None
    article_reference: str = ""
    section_reference: str = ""


@dataclass(frozen=True)
class LoadedDocument:
    """Normalized document content and source metadata."""

    source_file: str
    source_type: str
    title: str
    blocks: list[DocumentBlock]
    document_type: str = ""
    metadata: dict[str, str] = field(default_factory=dict)


def classify_blocks(blocks: list[DocumentBlock]) -> tuple[str, dict[str, str]]:
    """Classify a loaded document from its extracted text blocks."""

    text = "\n".join(block.text for block in blocks)
    document_type, scores = classify_document(text)
    return document_type, {
        "classification_law_score": str(scores["law_score"]),
        "classification_decision_score": str(scores["decision_score"]),
    }


def detect_legal_references(text: str) -> tuple[str, str]:
    """Return an article reference and broader legal section reference."""

    article_match = ARTICLE_PATTERN.match(text)
    section_match = LEGAL_SECTION_PATTERN.match(text)
    if not article_match:
        article_match = re.match(
            r"^\s*(?P<label>\u0627\u0644\u0645\u0627\u062f\u0629|\u0645\u0627\u062f\u0629)\s*"
            r"(?P<number>[\(\[]?[0-9\u0660-\u0669]+[\)\]]?)",
            text,
        )
    if not section_match:
        section_match = re.match(
            r"^\s*(?P<label>\u0627\u0644\u0628\u0627\u0628|\u0627\u0644\u0641\u0635\u0644|\u0627\u0644\u0642\u0633\u0645|\u0627\u0644\u0641\u0631\u0639)\s+"
            r"(?P<name>[^\n:]{1,120})",
            text,
        )
    article = article_match.group(0).strip() if article_match else ""
    section = section_match.group(0).strip() if section_match else ""
    return article, section


def _serialize_property(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip() if value is not None else ""


class DocumentLoader(ABC):
    """Interface implemented by each supported source format."""

    extensions: tuple[str, ...] = ()

    @abstractmethod
    def load(self, path: Path) -> LoadedDocument:
        """Load one source document."""


class DocxDocumentLoader(DocumentLoader):
    """Load Word documents while preserving paragraph and table order."""

    extensions = (".docx",)

    @staticmethod
    def _iter_body_items(document: Any) -> Iterable[Any]:
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                yield Paragraph(child, document)
            elif child.tag.endswith("}tbl"):
                yield Table(child, document)

    @staticmethod
    def _heading_level(paragraph: Any) -> int | None:
        style_name = paragraph.style.name if paragraph.style else ""
        style_match = HEADING_STYLE_PATTERN.match(style_name)
        if style_match:
            return int(style_match.group(1) or 1)

        outline = paragraph._p.xpath("./w:pPr/w:outlineLvl")
        if outline:
            value = outline[0].get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
            )
            if value is not None and value.isdigit():
                return int(value) + 1
        return None

    @staticmethod
    def _table_text(table: Any) -> str:
        rows = []
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        return "\n".join(rows)

    def load(self, path: Path) -> LoadedDocument:
        try:
            from docx import Document
            from docx.text.paragraph import Paragraph
        except ImportError as error:
            raise RuntimeError(
                "DOCX ingestion requires python-docx. "
                "Run: python -m pip install -r requirements.txt"
            ) from error

        document = Document(path)
        properties = document.core_properties
        metadata = {
            key: value
            for key, value in {
                "author": _serialize_property(properties.author),
                "subject": _serialize_property(properties.subject),
                "keywords": _serialize_property(properties.keywords),
                "category": _serialize_property(properties.category),
                "comments": _serialize_property(properties.comments),
                "created": _serialize_property(properties.created),
                "modified": _serialize_property(properties.modified),
                "last_modified_by": _serialize_property(
                    properties.last_modified_by
                ),
            }.items()
            if value
        }

        blocks = []
        first_heading = ""
        for item in self._iter_body_items(document):
            if isinstance(item, Paragraph):
                text = clean_text(item.text)
                if not text:
                    continue
                heading_level = self._heading_level(item)
                block_type = "heading" if heading_level else "paragraph"
                if heading_level and not first_heading:
                    first_heading = text
            else:
                text = self._table_text(item)
                if not text:
                    continue
                heading_level = None
                block_type = "table"

            article, section = detect_legal_references(text)
            blocks.append(
                DocumentBlock(
                    text=text,
                    block_type=block_type,
                    heading_level=heading_level,
                    article_reference=article,
                    section_reference=section,
                )
            )

        document_type, classification_metadata = classify_blocks(blocks)
        metadata.update(classification_metadata)
        title = clean_text(properties.title or "") or first_heading or path.stem
        return LoadedDocument(
            source_file=path.name,
            source_type="docx",
            title=title,
            blocks=blocks,
            document_type=document_type,
            metadata=metadata,
        )


class TxtDocumentLoader(DocumentLoader):
    """Load UTF-8 text files, including legacy page-separated exports."""

    extensions = (".txt",)

    def load(self, path: Path) -> LoadedDocument:
        content = clean_text(path.read_text(encoding="utf-8-sig"))
        separators = list(PAGE_SEPARATOR_PATTERN.finditer(content))
        pages: list[tuple[int, str]] = []

        if separators:
            for index, separator in enumerate(separators):
                end = (
                    separators[index + 1].start()
                    if index + 1 < len(separators)
                    else len(content)
                )
                pages.append(
                    (
                        int(separator.group(1)),
                        content[separator.end() : end].strip(),
                    )
                )
        else:
            pages = [(1, content)]

        blocks = []
        for page_number, page_text in pages:
            for paragraph in re.split(r"\n\s*\n|\n(?=\s*(?:Ø§Ù„Ù…Ø§Ø¯Ø©|Ø§Ù„Ø¨Ø§Ø¨|Ø§Ù„ÙØµÙ„|Ø§Ù„Ù‚Ø³Ù…|Ø§Ù„ÙØ±Ø¹)\b)", page_text):
                text = clean_text(paragraph)
                if not text:
                    continue
                article, section = detect_legal_references(text)
                blocks.append(
                    DocumentBlock(
                        text=text,
                        page_number=page_number,
                        article_reference=article,
                        section_reference=section,
                    )
                )

        document_type, metadata = classify_blocks(blocks)
        return LoadedDocument(
            source_file=path.name,
            source_type="txt",
            title=path.stem,
            blocks=blocks,
            document_type=document_type,
            metadata=metadata,
        )


class PdfDocumentLoader(DocumentLoader):
    """Load PDFs with text extraction first and OCR fallback for scanned pages."""

    extensions = (".pdf",)
    minimum_text_chars = 20
    ocr_language = "ara+eng"
    ocr_dpi = 400
    ocr_configs = ("--oem 1 --psm 6", "--oem 1 --psm 4", "--oem 1 --psm 11")

    @staticmethod
    def _ocr_quality_score(text: str) -> float:
        """Score OCR output by Arabic/legal signal and obvious noise."""

        if not text:
            return 0.0

        arabic_chars = len(re.findall(r"[\u0600-\u06ff]", text))
        digit_chars = len(re.findall(r"[0-9\u0660-\u0669]", text))
        legal_hits = len(
            re.findall(
                r"Ù‚Ø§Ù†ÙˆÙ†|Ù‚Ø±Ø§Ø±|Ø§Ù„Ù…Ø§Ø¯Ø©|Ù…Ø¬Ù„Ø³|Ø§Ù„Ù†ÙˆØ§Ø¨|Ø±Ù‚Ù…|Ù„Ø³Ù†Ø©",
                text,
            )
        )
        latin_noise = len(re.findall(r"[A-Za-z]{3,}", text))
        replacement_noise = text.count("?") + text.count("ï¿½")
        return (
            arabic_chars
            + (2.0 * digit_chars)
            + (20.0 * legal_hits)
            - (4.0 * latin_noise)
            - (10.0 * replacement_noise)
        )

    @staticmethod
    def _prepare_ocr_images(image: Any) -> list[Any]:
        """Create OCR variants for scanned Arabic PDFs."""

        from PIL import ImageFilter, ImageOps

        grayscale = ImageOps.grayscale(image)
        autocontrast = ImageOps.autocontrast(grayscale)
        sharpened = autocontrast.filter(ImageFilter.SHARPEN)
        threshold = sharpened.point(lambda value: 255 if value > 170 else 0)
        return [sharpened, threshold]

    def _ocr_page(self, path: Path, page_number: int) -> str:
        """OCR one 1-based PDF page using local Tesseract tooling."""

        try:
            import pypdfium2 as pdfium
            import pytesseract
        except ImportError as error:
            raise RuntimeError(
                "Scanned PDF ingestion requires OCR dependencies. "
                "Install requirements.txt, then install the Tesseract OCR "
                "engine with Arabic language data."
            ) from error

        try:
            pdf = pdfium.PdfDocument(str(path))
            page = pdf[page_number - 1]
            image = page.render(scale=self.ocr_dpi / 72).to_pil()
            candidates = []
            for prepared_image in self._prepare_ocr_images(image):
                for config in self.ocr_configs:
                    text = clean_text(
                        pytesseract.image_to_string(
                            prepared_image,
                            lang=self.ocr_language,
                            config=config,
                        )
                    )
                    candidates.append(
                        (self._ocr_quality_score(text), text)
                    )
            return max(candidates, default=(0.0, ""), key=lambda item: item[0])[1]
        except pytesseract.TesseractNotFoundError as error:
            raise RuntimeError(
                "Tesseract OCR executable was not found. Install Tesseract "
                "and make sure it is available on PATH."
            ) from error
        except pytesseract.TesseractError as error:
            raise RuntimeError(
                "Tesseract OCR failed. Verify Arabic language data is installed."
            ) from error

    def _extract_page_text(self, page: Any, path: Path, page_number: int) -> tuple[str, bool]:
        """Return page text and whether OCR was used."""

        text_layer = clean_text(page.extract_text() or "")
        if len(text_layer) >= self.minimum_text_chars:
            return text_layer, False

        ocr_text = self._ocr_page(path, page_number)
        if ocr_text:
            return ocr_text, True
        return text_layer, False

    def load(self, path: Path) -> LoadedDocument:
        try:
            import pdfplumber
        except ImportError as error:
            raise RuntimeError(
                "PDF ingestion requires pdfplumber. "
                "Run: python -m pip install -r requirements.txt"
            ) from error

        blocks = []
        used_ocr = False
        with pdfplumber.open(path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                page_text, page_used_ocr = self._extract_page_text(
                    page,
                    path,
                    page_index,
                )
                used_ocr = used_ocr or page_used_ocr
                if not page_text:
                    continue
                for paragraph in re.split(
                    r"\n\s*\n|\n(?=\s*(?:Ã˜Â§Ã™â€žÃ™â€¦Ã˜Â§Ã˜Â¯Ã˜Â©|Ã˜Â§Ã™â€žÃ˜Â¨Ã˜Â§Ã˜Â¨|Ã˜Â§Ã™â€žÃ™ÂÃ˜ÂµÃ™â€ž|Ã˜Â§Ã™â€žÃ™â€šÃ˜Â³Ã™â€¦|Ã˜Â§Ã™â€žÃ™ÂÃ˜Â±Ã˜Â¹)\b)",
                    page_text,
                ):
                    text = clean_text(paragraph)
                    if not text:
                        continue
                    article, section = detect_legal_references(text)
                    blocks.append(
                        DocumentBlock(
                            text=text,
                            page_number=page_index,
                            article_reference=article,
                            section_reference=section,
                        )
                    )

        document_type, metadata = classify_blocks(blocks)
        metadata["ocr"] = "tesseract" if used_ocr else "none"
        return LoadedDocument(
            source_file=path.name,
            source_type="pdf",
            title=path.stem,
            blocks=blocks,
            document_type=document_type,
            metadata=metadata,
        )


LOADERS = (DocxDocumentLoader(), TxtDocumentLoader(), PdfDocumentLoader())
LOADER_BY_EXTENSION = {
    extension: loader
    for loader in LOADERS
    for extension in loader.extensions
}


def supported_extensions() -> tuple[str, ...]:
    """Return source extensions currently supported by the registry."""

    return tuple(sorted(LOADER_BY_EXTENSION))


def load_document(path: Path) -> LoadedDocument:
    """Load one document with the registered loader for its extension."""

    loader = LOADER_BY_EXTENSION.get(path.suffix.lower())
    if loader is None:
        supported = ", ".join(supported_extensions())
        raise ValueError(
            f"Unsupported document type '{path.suffix}'. Supported: {supported}"
        )
    return loader.load(path)


def load_documents(folder: Path) -> list[LoadedDocument]:
    """Batch-load all supported documents in a folder."""

    extensions = set(supported_extensions())
    paths = sorted(
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() in extensions
        and not path.name.startswith("~$")
    )
    return [load_document(path) for path in paths]
