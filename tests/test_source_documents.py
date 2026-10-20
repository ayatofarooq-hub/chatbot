"""Verify original source downloads and offline previews without touching legal data."""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from urllib.parse import quote

from docx import Document
from starlette.testclient import TestClient

from app.api import app
from app.source_documents import original_word_file


class SourceDocumentTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.filename = "وثيقة رسمية.docx"
        self.path = self.root / self.filename
        document = Document()
        document.add_paragraph("First paragraph <script>alert('test')</script>")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Table value"
        table.cell(0, 1).text = "Second cell"
        document.add_paragraph("Final paragraph")
        document.save(self.path)
        self.registry = {"by_chunk_id": {"cited-chunk": {"source_file": self.filename}}}
        for target, value in [
            ("app.source_documents.WORD_ROOTS", (self.root,)),
            ("app.api.load_registry", self.registry),
            ("app.api.authenticated_admin", {"id": 1}),
        ]:
            patcher = patch(target, return_value=value) if target.startswith("app.api") else patch(target, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.client = TestClient(app)
        self.base = "/api/source-documents/" + quote(self.filename, safe="")

    def test_links_and_download_return_the_exact_original(self):
        info = self.client.get(self.base)
        self.assertEqual(info.status_code, 200)
        self.assertEqual(info.json()["filename"], self.filename)
        downloaded = self.client.get(info.json()["download_url"])
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.content, self.path.read_bytes())
        self.assertIn("attachment;", downloaded.headers["content-disposition"])
        self.assertIn("filename*=utf-8''", downloaded.headers["content-disposition"].lower())
        self.assertEqual(downloaded.headers["cache-control"], "no-store")

    def test_preview_preserves_block_order_and_escapes_source_content(self):
        response = self.client.get(self.base + "/preview")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/html", response.headers["content-type"])
        self.assertIn("&lt;script&gt;", response.text)
        self.assertNotIn("<script>", response.text)
        self.assertLess(response.text.index("First paragraph"), response.text.index("Table value"))
        self.assertLess(response.text.index("Table value"), response.text.index("Final paragraph"))

    def test_all_routes_require_authentication(self):
        with patch("app.api.authenticated_admin", return_value=None):
            for suffix in ["", "/preview", "/download"]:
                self.assertEqual(self.client.get(self.base + suffix).status_code, 401)

    def test_existing_but_unregistered_file_cannot_be_served(self):
        with patch("app.api.load_registry", return_value={"by_chunk_id": {}}):
            self.assertEqual(self.client.get(self.base + "/download").status_code, 404)

    def test_registered_but_missing_original_has_no_download_link(self):
        self.path.unlink()
        self.assertEqual(self.client.get(self.base).status_code, 404)

    def test_path_traversal_absolute_paths_and_non_word_files_are_rejected(self):
        for filename in ["../secret.docx", "..\\secret.docx", "C:\\secret.docx", "/secret.docx", "metadata.json"]:
            registry = {"by_chunk_id": {"test": {"source_file": filename}}}
            with self.subTest(filename=filename), self.assertRaises(LookupError):
                original_word_file(filename, registry)

    def test_ambiguous_original_is_not_guessed(self):
        other = self.root / "duplicate"
        other.mkdir()
        (other / self.filename).write_bytes(self.path.read_bytes())
        with patch("app.source_documents.WORD_ROOTS", (self.root, other)):
            self.assertEqual(self.client.get(self.base).status_code, 404)

    def test_invalid_word_file_can_be_downloaded_without_crashing_preview(self):
        self.path.write_bytes(b"unreadable Word file")
        self.assertEqual(self.client.get(self.base + "/preview").status_code, 422)
        self.assertEqual(self.client.get(self.base + "/download").content, self.path.read_bytes())


if __name__ == "__main__":
    unittest.main()
