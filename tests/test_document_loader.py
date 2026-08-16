"""Unit tests for supported document loading formats."""

import tempfile
import unittest
from pathlib import Path

from app.extraction.document_loader import LoadedDocument, load_document


class DocumentLoaderTests(unittest.TestCase):
    def test_loads_plain_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.txt"
            path.write_text("This is a plain text file.", encoding="utf-8")
            document = load_document(path)

        self.assertIsInstance(document, LoadedDocument)
        self.assertEqual(document.source_file, "sample.txt")
        self.assertEqual(document.filename, "sample.txt")
        self.assertEqual(document.extension, ".txt")
        self.assertEqual(document.text, "This is a plain text file.")
        self.assertEqual(document.metadata, {})

    def test_loads_markdown(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.md"
            path.write_text("# Heading\n\nThis is markdown.", encoding="utf-8")
            document = load_document(path)

        self.assertEqual(document.text, "# Heading\n\nThis is markdown.")
        self.assertEqual(document.filename, "sample.md")
        self.assertEqual(document.extension, ".md")
        self.assertEqual(document.metadata, {})

    def test_loads_html(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.html"
            path.write_text(
                "<html><body><h1>Title</h1><p>Paragraph text.</p></body></html>",
                encoding="utf-8",
            )
            document = load_document(path)

        self.assertIn("Title", document.text)
        self.assertIn("Paragraph text.", document.text)

    def test_loads_rtf(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.rtf"
            path.write_text(
                "{\\rtf1\\ansi{\\b Bold text}\\par Normal text.}",
                encoding="utf-8",
            )
            document = load_document(path)

        self.assertIn("Bold text", document.text)
        self.assertIn("Normal text.", document.text)

    def test_loads_docx(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.docx"
            document = Document()
            document.add_paragraph("First paragraph.")
            document.add_paragraph("Second paragraph.")
            document.save(path)

            loaded = load_document(path)

        self.assertIn("First paragraph.", loaded.text)
        self.assertIn("Second paragraph.", loaded.text)

    def test_loads_pdf(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.pdf"
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "PDF sample text.")
            pdf.save(path)
            pdf.close()

            loaded = load_document(path)

        self.assertIn("PDF sample text.", loaded.text)

    def test_unsupported_suffix_raises(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.exe"
            path.write_text("binary content", encoding="utf-8")

            with self.assertRaises(ValueError):
                load_document(path)


if __name__ == "__main__":
    unittest.main()
